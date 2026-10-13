import sys
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml

# Color Palette
PRIMARY_BLUE = RGBColor(0, 112, 192)      # #0070C0 (Template Banner Blue)
DARK_NAVY    = RGBColor(0, 51, 102)       # #003366 (Deep Navy Title)
ACCENT_BLUE  = RGBColor(30, 136, 229)     # #1E88E5 (Accent Blue)
BG_LIGHT     = RGBColor(248, 250, 252)    # #F8FAFC (Slate light)
BORDER_GRAY  = RGBColor(203, 213, 225)    # #CBD5E1 (Card border)
TEXT_DARK    = RGBColor(31, 41, 55)       # #1F2937 (Charcoal header)
TEXT_MUTED   = RGBColor(75, 85, 99)       # #4B5563 (Body muted)
WHITE        = RGBColor(255, 255, 255)

# Status Colors
GREEN_BG     = RGBColor(220, 252, 231)    # #DCFCE7
GREEN_BORDER = RGBColor(22, 163, 74)      # #16A34A
GREEN_TEXT   = RGBColor(21, 128, 61)      # #15803D

AMBER_BG     = RGBColor(254, 243, 199)    # #FEF3C7
AMBER_BORDER = RGBColor(217, 119, 6)      # #D97706
AMBER_TEXT   = RGBColor(180, 83, 9)       # #B45309

RED_BG       = RGBColor(254, 226, 226)    # #FEE2E2
RED_BORDER   = RGBColor(220, 38, 38)      # #DC2626
RED_TEXT     = RGBColor(185, 28, 28)      # #B91C1C

GRAY_BG      = RGBColor(243, 244, 246)    # #F3F4F6
GRAY_BORDER  = RGBColor(156, 163, 175)    # #9CA3AF
GRAY_TEXT    = RGBColor(75, 85, 99)       # #4B5563

def strip_bullet(paragraph):
    """Explicitly removes any bullet formatting from a paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    for child in list(pPr):
        tag = child.tag.split('}')[-1]
        if tag.startswith('bu'):
            pPr.remove(child)
    pPr.append(parse_xml('<a:buNone xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'))

def add_card(slide, left, top, width, height, bg_color=BG_LIGHT, border_color=BORDER_GRAY, border_width=1):
    """Creates a stylized rounded card container."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(border_width)
    else:
        shape.line.fill.background()
    return shape

def add_arrow(slide, left, top, width, height, direction='right', color=PRIMARY_BLUE):
    """Creates a directional connector arrow."""
    stype = MSO_SHAPE.RIGHT_ARROW if direction == 'right' else MSO_SHAPE.DOWN_ARROW
    shape = slide.shapes.add_shape(stype, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_pill(slide, left, top, width, height, text, bg_color, border_color, text_color, font_size=10, bold=True):
    """Creates a small status pill/badge."""
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    pill.fill.solid()
    pill.fill.fore_color.rgb = bg_color
    pill.line.color.rgb = border_color
    pill.line.width = Pt(1)
    tf = pill.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    strip_bullet(p)
    r = p.add_run()
    r.text = text
    r.font.name = "Arial"
    r.font.size = Pt(font_size)
    r.font.bold = bold
    r.font.color.rgb = text_color
    return pill

def build_presentation():
    src_path = os.path.abspath('ppt/IDE BOOTCAMP PPT FORMAT.pptx')
    out_pptx_path = os.path.abspath('ppt/JIVA_IDE_BOOTCAMP_FORMAT_FINAL.pptx')
    out_pdf_path = os.path.abspath('ppt/JIVA_IDE_BOOTCAMP_FORMAT_FINAL.pdf')

    print(f"Loading template from: {src_path}")
    prs = Presentation(src_path)

    # =========================================================================
    # SLIDE 1: TITLE PAGE (ENHANCED CARD METADATA)
    # =========================================================================
    s1 = prs.slides[0]
    for shape in s1.shapes:
        if shape.name == 'Title 7':
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "IDE BOOTCAMP 2026"
            p.font.name = "Garamond"
            p.font.size = Pt(40)
            p.font.bold = True
            p.font.color.rgb = DARK_NAVY
            
        elif shape.name == 'Subtitle 3':
            shape.left = Inches(0.80)
            shape.width = Inches(11.73)
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "JIVA: REAL-TIME HEALTHCARE COORDINATION"
            p.font.name = "Times New Roman"
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = PRIMARY_BLUE
            
        elif shape.name == 'TextBox 9':
            # Hide old textbox
            shape.left = Inches(0)
            shape.top = Inches(0)
            shape.width = Inches(0.1)
            shape.height = Inches(0.1)
            shape.text_frame.clear()

    # Add stylish container card for Slide 1 metadata
    m_card = add_card(s1, Inches(0.60), Inches(2.10), Inches(5.60), Inches(4.75), bg_color=WHITE, border_color=PRIMARY_BLUE, border_width=1.5)
    mtf = m_card.text_frame
    mtf.margin_left = Inches(0.15)
    mtf.margin_right = Inches(0.15)
    mtf.margin_top = Inches(0.12)
    mtf.margin_bottom = Inches(0.10)

    mp0 = mtf.paragraphs[0]
    strip_bullet(mp0)
    mr0 = mp0.add_run()
    mr0.text = "PROJECT IDENTIFICATION & CORE THESIS\n"
    mr0.font.name = "Arial"
    mr0.font.size = Pt(11.5)
    mr0.font.bold = True
    mr0.font.color.rgb = DARK_NAVY

    items = [
        ("Problem Statement Title: ", "Real-Time Emergency Healthcare Coordination & Dynamic Hospital Routing"),
        ("Theme: ", "Healthcare & MedTech / Cloud-Native Emergency Mesh"),
        ("PS Category: ", "Software (Event-Driven Healthcare Mesh)"),
        ("Team Name: ", "Monochrome"),
        ("Team Members: ", "Yukith M Joseph  •  Pranathi Nagabhushana"),
        ("Solution Summary: ", "An event-driven layer maintaining a shared, timestamped care state for patients, ambulances, and hospitals, routing via explicit hospital acceptance and dynamically rerouting when state changes."),
        ("Verification Status: ", "Fully functional local prototype verified on synthetic Bengaluru data (124/124 checks pass). AWS architecture modeled in CDK (not deployed).")
    ]

    for label, val in items:
        p = mtf.add_paragraph()
        strip_bullet(p)
        p.space_before = Pt(4)
        p.line_spacing = 1.15
        
        r_lbl = p.add_run()
        r_lbl.text = "• " + label
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(9.5)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = PRIMARY_BLUE

        r_val = p.add_run()
        r_val.text = val
        r_val.font.name = "Arial"
        r_val.font.size = Pt(9)
        r_val.font.bold = False
        r_val.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 2: IDEA TITLE (CLOSED-LOOP COORDINATION & STATUS BADGES)
    # =========================================================================
    s2 = prs.slides[1]
    for s in s2.shapes:
        if s.name == 'Title 1':
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "JIVA: CLOSED-LOOP HEALTHCARE COORDINATION"
            p.font.name = "Times New Roman"
            p.font.size = Pt(32)
            p.font.bold = True
            p.font.color.rgb = DARK_NAVY
        elif s.name == 'TextBox 8':
            s.left = Inches(0)
            s.top = Inches(0)
            s.width = Inches(0.1)
            s.height = Inches(0.1)
            s.text_frame.clear()

    # --- TOP: Closed-Loop Decision Flowchart ---
    add_card(s2, Inches(0.60), Inches(1.25), Inches(12.13), Inches(2.55), bg_color=WHITE, border_color=PRIMARY_BLUE, border_width=1.5)

    banner_tb = s2.shapes.add_textbox(Inches(0.80), Inches(1.35), Inches(11.70), Inches(0.35))
    tf = banner_tb.text_frame
    p = tf.paragraphs[0]
    p.text = "PROPOSED SOLUTION: The Closed-Loop Hospital Acceptance & Dynamic Rerouting Engine"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY

    # 5 Flowchart Step Boxes
    steps = [
        ("1. Emergency Intake", "Patient / Paramedic logs incident\nClinical requirements tagged\nCanonical event ingested", Inches(0.80)),
        ("2. Feasibility Gate", "Candidate hospitals evaluated\nDeterminism & freshness check\nIdentifies eligible facilities", Inches(3.15)),
        ("3. Acceptance SLA", "Targeted request dispatched\nCase-scoped confirmation query\nStrict timeout countdown", Inches(5.50)),
        ("4. Explicit Response", "Mandatory active confirmation:\n[ACCEPTED]  [LIMITED]\n[REJECTED]  [UNAVAILABLE]\nSilence expires to UNKNOWN", Inches(7.85)),
        ("5. Dynamic Reroute", "Confirmed unit dispatched\nReal-time GPS telemetry sync\nReroutes if hospital degrades", Inches(10.20))
    ]

    for title, desc, l_pos in steps:
        box = add_card(s2, l_pos, Inches(1.75), Inches(2.05), Inches(1.85), bg_color=BG_LIGHT, border_color=BORDER_GRAY)
        btf = box.text_frame
        btf.margin_left = Inches(0.08)
        btf.margin_right = Inches(0.08)
        btf.margin_top = Inches(0.08)
        btf.margin_bottom = Inches(0.08)
        bp = btf.paragraphs[0]
        strip_bullet(bp)
        br1 = bp.add_run()
        br1.text = title + "\n"
        br1.font.name = "Arial"
        br1.font.size = Pt(11)
        br1.font.bold = True
        br1.font.color.rgb = PRIMARY_BLUE
        
        br2 = bp.add_run()
        br2.text = desc
        br2.font.name = "Arial"
        br2.font.size = Pt(9.5)
        br2.font.color.rgb = TEXT_MUTED

    # Connecting Arrows
    add_arrow(s2, Inches(2.90), Inches(2.55), Inches(0.20), Inches(0.22))
    add_arrow(s2, Inches(5.25), Inches(2.55), Inches(0.20), Inches(0.22))
    add_arrow(s2, Inches(7.60), Inches(2.55), Inches(0.20), Inches(0.22))
    add_arrow(s2, Inches(9.95), Inches(2.55), Inches(0.20), Inches(0.22))

    # --- BOTTOM: Side-by-Side Architectural Contrast Cards ---
    # Left: Existing Approach
    c_exist = add_card(s2, Inches(0.60), Inches(3.95), Inches(5.95), Inches(2.80), bg_color=WHITE, border_color=RED_BORDER, border_width=1.5)
    ctf = c_exist.text_frame
    ctf.margin_left = Inches(0.12)
    ctf.margin_right = Inches(0.12)
    ctf.margin_top = Inches(0.10)
    cp = ctf.paragraphs[0]
    strip_bullet(cp)
    cr = cp.add_run()
    cr.text = "❌  EXISTING REALITY: Static & Unverified Assumptions\n"
    cr.font.name = "Arial"
    cr.font.size = Pt(12.5)
    cr.font.bold = True
    cr.font.color.rgb = RED_TEXT

    exist_bullets = [
        ("• Assumed Readiness: ", "Directory listings prove a hospital exists, not that an ICU bed or specialist is free right now."),
        ("• Fragmented Phone Calls: ", "Paramedics negotiate en-route over phone; no shared timestamped picture."),
        ("• Fatal ER Diversion: ", "Ambulance arrives unannounced at saturated emergency room → turned away at the door."),
        ("• Static Destination Lock: ", "Destination is chosen once at dispatch; zero awareness if conditions degrade en-route.")
    ]
    for b_title, b_desc in exist_bullets:
        bp = ctf.add_paragraph()
        strip_bullet(bp)
        bp.space_before = Pt(3)
        bp.line_spacing = 1.15
        r1 = bp.add_run()
        r1.text = b_title
        r1.font.name = "Arial"
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = TEXT_DARK
        r2 = bp.add_run()
        r2.text = b_desc
        r2.font.name = "Arial"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_MUTED

    # Right: JIVA Approach
    c_jiva = add_card(s2, Inches(6.78), Inches(3.95), Inches(5.95), Inches(2.80), bg_color=WHITE, border_color=GREEN_BORDER, border_width=1.5)
    jtf = c_jiva.text_frame
    jtf.margin_left = Inches(0.12)
    jtf.margin_right = Inches(0.12)
    jtf.margin_top = Inches(0.10)
    jp = jtf.paragraphs[0]
    strip_bullet(jp)
    jr = jp.add_run()
    jr.text = "✔  JIVA SOLUTION: Closed-Loop Event-Driven Coordination\n"
    jr.font.name = "Arial"
    jr.font.size = Pt(12.5)
    jr.font.bold = True
    jr.font.color.rgb = GREEN_TEXT

    jiva_bullets = [
        ("• Explicit Confirmation: ", "Hospitals must actively confirm acceptance; unconfirmed requests deterministically expire to UNKNOWN."),
        ("• Live Dynamic Rerouting: ", "If an accepted hospital becomes UNAVAILABLE en-route, JIVA immediately triggers dynamic rerouting."),
        ("• Zero EHR Intrusion: ", "Lightweight web overlay across 4 role-scoped apps (:5174-:5177); zero costly HIS integrations."),
        ("• Deterministic Feasibility: ", "100% deterministic rule verification & SHA-256 event audits; generative AI strictly advisory.")
    ]
    for b_title, b_desc in jiva_bullets:
        bp = jtf.add_paragraph()
        strip_bullet(bp)
        bp.space_before = Pt(3)
        bp.line_spacing = 1.15
        r1 = bp.add_run()
        r1.text = b_title
        r1.font.name = "Arial"
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = TEXT_DARK
        r2 = bp.add_run()
        r2.text = b_desc
        r2.font.name = "Arial"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 3: TECHNICAL APPROACH (ARCHITECTURE DIAGRAM & PIPELINE)
    # =========================================================================
    s3 = prs.slides[2]
    for s in s3.shapes:
        if s.name == 'Title 1':
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "TECHNICAL APPROACH & ARCHITECTURE"
            p.font.name = "Times New Roman"
            p.font.size = Pt(32)
            p.font.bold = True
            p.font.color.rgb = DARK_NAVY
        elif s.name == 'TextBox 8':
            s.left = Inches(0)
            s.top = Inches(0)
            s.width = Inches(0.1)
            s.height = Inches(0.1)
            s.text_frame.clear()

    # --- TOP: Full-Stack Multi-Tier Architecture Diagram ---
    add_card(s3, Inches(0.60), Inches(1.25), Inches(12.13), Inches(3.40), bg_color=WHITE, border_color=PRIMARY_BLUE, border_width=1.5)

    arch_title = s3.shapes.add_textbox(Inches(0.80), Inches(1.30), Inches(11.70), Inches(0.30))
    p = arch_title.text_frame.paragraphs[0]
    p.text = "SYSTEM ARCHITECTURE: Hybrid Local-to-Cloud Event-Driven Mesh"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY

    # Tier 1: Clients & Roles
    t1 = add_card(s3, Inches(0.80), Inches(1.70), Inches(2.20), Inches(2.75), bg_color=BG_LIGHT, border_color=BORDER_GRAY)
    tf = t1.text_frame
    tf.margin_top = Inches(0.08)
    p = tf.paragraphs[0]
    strip_bullet(p)
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "CLIENT TIER\nReact 19 / Vite"
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = PRIMARY_BLUE

    client_pills = [
        ("Patient Portal", ":5174"),
        ("Paramedic App", ":5175"),
        ("Hospital ED", ":5176"),
        ("Coordinator", ":5177")
    ]
    for cname, cport in client_pills:
        cp = tf.add_paragraph()
        strip_bullet(cp)
        cp.space_before = Pt(4)
        cp.alignment = PP_ALIGN.CENTER
        r1 = cp.add_run()
        r1.text = f"[{cname} {cport}]"
        r1.font.name = "Arial"
        r1.font.size = Pt(9.5)
        r1.font.bold = True
        r1.font.color.rgb = TEXT_DARK

    add_arrow(s3, Inches(3.05), Inches(2.95), Inches(0.20), Inches(0.22))

    # Tier 2: Ingestion & Auth Gate
    t2 = add_card(s3, Inches(3.30), Inches(1.70), Inches(2.20), Inches(2.75), bg_color=BG_LIGHT, border_color=BORDER_GRAY)
    tf = t2.text_frame
    tf.margin_top = Inches(0.08)
    p = tf.paragraphs[0]
    strip_bullet(p)
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "INGESTION & AUTH\nAPI Gateway / REST"
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = PRIMARY_BLUE

    gate_items = [
        "Zod Schema Validation",
        "Strict Payload Filtering",
        "Cognito / Role Tokens",
        "Policy Hash (hashPolicy)",
        "Port-Keyed Isolation"
    ]
    for gitem in gate_items:
        gp = tf.add_paragraph()
        strip_bullet(gp)
        gp.space_before = Pt(3)
        gp.alignment = PP_ALIGN.CENTER
        r = gp.add_run()
        r.text = "• " + gitem
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.color.rgb = TEXT_MUTED

    add_arrow(s3, Inches(5.55), Inches(2.95), Inches(0.20), Inches(0.22))

    # Tier 3: Event Engine & Microservices
    t3 = add_card(s3, Inches(5.80), Inches(1.70), Inches(3.60), Inches(2.75), bg_color=BG_LIGHT, border_color=BORDER_GRAY)
    tf = t3.text_frame
    tf.margin_top = Inches(0.08)
    p = tf.paragraphs[0]
    strip_bullet(p)
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "EVENT MESH & DOMAIN LOGIC\nAmazon EventBridge + Lambda / Node.js"
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = PRIMARY_BLUE

    mesh_items = [
        ("Canonical Event Bus", "Broadcasts validated lifecycle events"),
        ("State Transition Engine", "Tracks PENDING → ACCEPTED / REJECTED"),
        ("Dynamic Reroute Logic", "Auto-triggers recalculation on status change"),
        ("Deterministic Shadow Engine", "Parallel feasibility, freshness & provenance check"),
        ("Amazon Bedrock (Advisory)", "Strictly clinical summaries; ZERO routing authority")
    ]
    for mtitle, mdesc in mesh_items:
        mp = tf.add_paragraph()
        strip_bullet(mp)
        mp.space_before = Pt(2)
        r1 = mp.add_run()
        r1.text = "• " + mtitle + ": "
        r1.font.name = "Arial"
        r1.font.size = Pt(9)
        r1.font.bold = True
        r1.font.color.rgb = TEXT_DARK
        r2 = mp.add_run()
        r2.text = mdesc
        r2.font.name = "Arial"
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = TEXT_MUTED

    add_arrow(s3, Inches(9.45), Inches(2.95), Inches(0.20), Inches(0.22))

    # Tier 4: Persistence & Spatial Engine
    t4 = add_card(s3, Inches(9.70), Inches(1.70), Inches(2.80), Inches(2.75), bg_color=BG_LIGHT, border_color=BORDER_GRAY)
    tf = t4.text_frame
    tf.margin_top = Inches(0.08)
    p = tf.paragraphs[0]
    strip_bullet(p)
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "PERSISTENCE & SPATIAL\nDynamoDB & MapLibre GL"
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = PRIMARY_BLUE

    persist_items = [
        ("Operational State Store", "In-memory locally (<50ms) / DynamoDB cloud"),
        ("SHA-256 Audit Ledger", "Immutable forensic history of every state change"),
        ("Spatial Map Engine", "MapLibre GL client-side WebGL vector rendering"),
        ("Routing Engine Fallback", "Valhalla/OSRM; failures yield UNKNOWN (not 0)")
    ]
    for ptitle, pdesc in persist_items:
        pp = tf.add_paragraph()
        strip_bullet(pp)
        pp.space_before = Pt(3)
        r1 = pp.add_run()
        r1.text = "• " + ptitle + ": "
        r1.font.name = "Arial"
        r1.font.size = Pt(9)
        r1.font.bold = True
        r1.font.color.rgb = TEXT_DARK
        r2 = pp.add_run()
        r2.text = pdesc
        r2.font.name = "Arial"
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = TEXT_MUTED

    # --- BOTTOM: Implementation Pipeline & Verified Proof Badges ---
    add_card(s3, Inches(0.60), Inches(4.80), Inches(12.13), Inches(1.95), bg_color=WHITE, border_color=BORDER_GRAY)

    pip_tb = s3.shapes.add_textbox(Inches(0.80), Inches(4.85), Inches(11.70), Inches(0.25))
    p = pip_tb.text_frame.paragraphs[0]
    p.text = "IMPLEMENTATION METHODOLOGY & AUTOMATED VERIFICATION BASELINE"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY

    psteps = [
        ("Phase 1: Ingestion", "Zod schemas, event bus, & role contracts"),
        ("Phase 2: Acceptance", "Case-scoped requests & timeout SLA"),
        ("Phase 3: Dynamic Reroute", "Auto-rerouting upon hospital degradation"),
        ("Phase 4: Multi-App Mesh", "4 role-scoped apps + Live Demo portal")
    ]
    for idx, (pname, pdetail) in enumerate(psteps):
        l_pos = Inches(0.80 + idx * 2.95)
        pbox = add_card(s3, l_pos, Inches(5.15), Inches(2.75), Inches(0.70), bg_color=BG_LIGHT, border_color=PRIMARY_BLUE)
        ptf = pbox.text_frame
        ptf.margin_top = Inches(0.04)
        pp = ptf.paragraphs[0]
        strip_bullet(pp)
        pr1 = pp.add_run()
        pr1.text = pname + "\n"
        pr1.font.name = "Arial"
        pr1.font.size = Pt(10)
        pr1.font.bold = True
        pr1.font.color.rgb = PRIMARY_BLUE
        pr2 = pp.add_run()
        pr2.text = pdetail
        pr2.font.name = "Arial"
        pr2.font.size = Pt(8.5)
        pr2.font.color.rgb = TEXT_MUTED

    add_pill(s3, Inches(0.80), Inches(6.00), Inches(3.70), Inches(0.60),
             "✔ 124/124 Runtime Checks Passing\nFlagship simulation & parity verified",
             GREEN_BG, GREEN_BORDER, GREEN_TEXT, font_size=9)

    add_pill(s3, Inches(4.80), Inches(6.00), Inches(3.70), Inches(0.60),
             "✔ 25 Unit Test Suites Verified\nDomain, transit, & auth coverage",
             GREEN_BG, GREEN_BORDER, GREEN_TEXT, font_size=9)

    add_pill(s3, Inches(8.80), Inches(6.00), Inches(3.90), Inches(0.60),
             "✔ 7 Lambda Bundles & CDK Synth\nesbuild compiled & cleanly synthesized",
             GREEN_BG, GREEN_BORDER, GREEN_TEXT, font_size=9)

    # =========================================================================
    # SLIDE 4: FEASIBILITY AND VIABILITY (PILLARS & RISK MITIGATION MATRIX)
    # =========================================================================
    s4 = prs.slides[3]
    for s in s4.shapes:
        if s.name == 'Title 1':
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "FEASIBILITY AND VIABILITY"
            p.font.name = "Times New Roman"
            p.font.size = Pt(36)
            p.font.bold = True
            p.font.color.rgb = DARK_NAVY
        elif s.name == 'TextBox 8':
            s.left = Inches(0)
            s.top = Inches(0)
            s.width = Inches(0.1)
            s.height = Inches(0.1)
            s.text_frame.clear()

    # --- TOP: 3 Feasibility Pillar Cards ---
    pillars = [
        ("Operational Feasibility", "Zero EHR Intrusion", [
            ("Lightweight Web Overlay: ", "Runs in any browser without costly integration into proprietary hospital EHR/HIS systems."),
            ("Zero Hardware Lock-In: ", "Compatible with any ambulance tablet, phone, or ER desktop; ECG IoT simulated via JSON."),
            ("Sub-50ms Processing: ", "Synchronous in-memory pipeline ensures instant operational response times.")
        ], Inches(0.60)),
        ("Technical Feasibility", "Offline-First to Cloud", [
            ("Self-Contained Local Prototype: ", "100% offline-capable on port-isolated Vite servers for field / hackathon resilience."),
            ("CDK Serverless Architecture: ", "Clean 1:1 mapping from local components to AWS EventBridge, Lambda, and DynamoDB."),
            ("Shadow Engine Verification: ", "Runs parallel deterministic checks with 0 regression disagreements.")
        ], Inches(4.75)),
        ("Economic Viability", "High-Impact Cost Efficiency", [
            ("Serverless Pay-Per-Event: ", "Negligible idle cost; scales dynamically with emergency call volume at fractions of a cent."),
            ("Prevents Fatal Diversions: ", "Eliminates catastrophic costs associated with secondary transfers and ER turnaways."),
            ("Resource Optimization: ", "Reduces wasted ambulance mileage and enables prepared hospital intake.")
        ], Inches(8.90))
    ]

    for ptitle, psub, pbullets, l_pos in pillars:
        pcard = add_card(s4, l_pos, Inches(1.25), Inches(3.85), Inches(2.20), bg_color=WHITE, border_color=PRIMARY_BLUE, border_width=1.5)
        ptf = pcard.text_frame
        ptf.margin_left = Inches(0.10)
        ptf.margin_right = Inches(0.10)
        ptf.margin_top = Inches(0.08)
        pp = ptf.paragraphs[0]
        strip_bullet(pp)
        pr1 = pp.add_run()
        pr1.text = ptitle + "\n"
        pr1.font.name = "Arial"
        pr1.font.size = Pt(12)
        pr1.font.bold = True
        pr1.font.color.rgb = DARK_NAVY

        pr2 = pp.add_run()
        pr2.text = psub + "\n"
        pr2.font.name = "Arial"
        pr2.font.size = Pt(9.5)
        pr2.font.bold = True
        pr2.font.color.rgb = PRIMARY_BLUE

        for b_name, b_val in pbullets:
            bp = ptf.add_paragraph()
            strip_bullet(bp)
            bp.space_before = Pt(2)
            bp.line_spacing = 1.15
            r1 = bp.add_run()
            r1.text = "• " + b_name
            r1.font.name = "Arial"
            r1.font.size = Pt(8.5)
            r1.font.bold = True
            r1.font.color.rgb = TEXT_DARK
            r2 = bp.add_run()
            r2.text = b_val
            r2.font.name = "Arial"
            r2.font.size = Pt(8.5)
            r2.font.color.rgb = TEXT_MUTED

    # --- BOTTOM: Visual Risk vs. Strategy Mitigation Matrix ---
    add_card(s4, Inches(0.60), Inches(3.60), Inches(12.13), Inches(3.15), bg_color=WHITE, border_color=BORDER_GRAY)

    rm_tb = s4.shapes.add_textbox(Inches(0.80), Inches(3.68), Inches(11.70), Inches(0.25))
    p = rm_tb.text_frame.paragraphs[0]
    p.text = "RISK VS. STRATEGY MITIGATION MATRIX"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY

    matrix_rows = [
        ("Public Directory Blind Spots", "Govt hospital registries list facilities but lack live ICU / bed status.",
         "Honest Availability Semantics", "Unconfirmed capacities strictly flagged UNKNOWN; explicit acceptance is mandatory.",
         Inches(4.00)),
        ("Urban Transit Cellular Drops", "Mobile connectivity fluctuates during high-speed ambulance transit.",
         "Offline Caching & Reconnection", "Local client state caching, auto Socket.IO backoff; keeps last GPS with UNKNOWN ETA.",
         Inches(4.65)),
        ("In-Memory State Volatility", "Prototype runs in memory; process crash resets active emergency session.",
         "Enterprise Persistence Path", "Migration roadmap mapped to transactional SQLite locally and DynamoDB single-table cloud.",
         Inches(5.30)),
        ("Generative AI Hallucinations", "Unbounded LLMs could hallucinate clinical triaging or routing destinations.",
         "Strict Advisory AI Sandboxing", "Amazon Bedrock strictly sandboxed to read-only summaries; routing is 100% deterministic code.",
         Inches(5.95))
    ]

    for r_title, r_desc, s_title, s_desc, top_pos in matrix_rows:
        rbox = add_card(s4, Inches(0.80), top_pos, Inches(5.60), Inches(0.58), bg_color=RED_BG, border_color=RED_BORDER)
        rtf = rbox.text_frame
        rtf.margin_left = Inches(0.08)
        rtf.margin_top = Inches(0.04)
        rp = rtf.paragraphs[0]
        strip_bullet(rp)
        rr1 = rp.add_run()
        rr1.text = "⚠ [RISK] " + r_title + ": "
        rr1.font.name = "Arial"
        rr1.font.size = Pt(9)
        rr1.font.bold = True
        rr1.font.color.rgb = RED_TEXT
        rr2 = rp.add_run()
        rr2.text = r_desc
        rr2.font.name = "Arial"
        rr2.font.size = Pt(8.5)
        rr2.font.color.rgb = TEXT_DARK

        add_arrow(s4, Inches(6.50), top_pos + Inches(0.18), Inches(0.18), Inches(0.20), color=PRIMARY_BLUE)

        sbox = add_card(s4, Inches(6.78), top_pos, Inches(5.75), Inches(0.58), bg_color=GREEN_BG, border_color=GREEN_BORDER)
        stf = sbox.text_frame
        stf.margin_left = Inches(0.08)
        stf.margin_top = Inches(0.04)
        sp = stf.paragraphs[0]
        strip_bullet(sp)
        sr1 = sp.add_run()
        sr1.text = "✔ [STRATEGY] " + s_title + ": "
        sr1.font.name = "Arial"
        sr1.font.size = Pt(9)
        sr1.font.bold = True
        sr1.font.color.rgb = GREEN_TEXT
        sr2 = sp.add_run()
        sr2.text = s_desc
        sr2.font.name = "Arial"
        sr2.font.size = Pt(8.5)
        sr2.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 5: IMPACT AND BENEFITS (4-STAKEHOLDER GRID & BENEFIT CALLOUTS)
    # =========================================================================
    s5 = prs.slides[4]
    for s in s5.shapes:
        if s.name == 'Title 1':
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "IMPACT AND BENEFITS"
            p.font.name = "Times New Roman"
            p.font.size = Pt(36)
            p.font.bold = True
            p.font.color.rgb = DARK_NAVY
        elif s.name == 'TextBox 8':
            s.left = Inches(0)
            s.top = Inches(0)
            s.width = Inches(0.1)
            s.height = Inches(0.1)
            s.text_frame.clear()

    # --- LEFT: 4-Quadrant Stakeholder Impact Grid ---
    stakeholders = [
        ("Patients & Families", "Transparent, Honest Emergency Transit", [
            "Honest pre-arrival travel ETAs (never fake 0-min ETAs).",
            "Eliminates traumatizing ER diversions & secondary hospital turnaways.",
            "Live transparent care status tracking during critical care intervals."
        ], Inches(0.60), Inches(1.25), PRIMARY_BLUE),

        ("Paramedic Crews", "Confirmed Destinations & Zero Phone Chaos", [
            "Guaranteed confirmed hospital bed before ambulance leaves the scene.",
            "Automated hands-free dynamic reroute alerts if receiving ER status degrades.",
            "Replaces chaotic en-route phone calls with structured digital handover."
        ], Inches(4.55), Inches(1.25), DARK_NAVY),

        ("Hospital Emergency Staff", "Structured Pre-Arrival Telemetry & Prep", [
            "Structured pre-arrival vitals enable advance trauma & cath lab preparation.",
            "One-click intake control: ACCEPT, LIMITED, or REJECT based on live capacity.",
            "Prevents ER bottleneck congestion and uncoordinated ambulance flooding."
        ], Inches(0.60), Inches(4.00), DARK_NAVY),

        ("Regional Care Coordinators", "District-Wide Situational Awareness", [
            "Real-time operational dashboard across regional hospital beds & dispatches.",
            "Active monitoring of acceptance requests & expiry SLA countdowns.",
            "Dynamic cross-facility load balancing during mass-casualty incidents."
        ], Inches(4.55), Inches(4.00), PRIMARY_BLUE)
    ]

    for stitle, ssub, sbullets, l_pos, t_pos, border_c in stakeholders:
        scard = add_card(s5, l_pos, t_pos, Inches(3.80), Inches(2.60), bg_color=WHITE, border_color=border_c, border_width=1.5)
        stf = scard.text_frame
        stf.margin_left = Inches(0.10)
        stf.margin_right = Inches(0.10)
        stf.margin_top = Inches(0.08)
        sp = stf.paragraphs[0]
        strip_bullet(sp)
        sr1 = sp.add_run()
        sr1.text = stitle + "\n"
        sr1.font.name = "Arial"
        sr1.font.size = Pt(11.5)
        sr1.font.bold = True
        sr1.font.color.rgb = DARK_NAVY

        sr2 = sp.add_run()
        sr2.text = ssub + "\n"
        sr2.font.name = "Arial"
        sr2.font.size = Pt(9)
        sr2.font.bold = True
        sr2.font.color.rgb = PRIMARY_BLUE

        for bitem in sbullets:
            bp = stf.add_paragraph()
            strip_bullet(bp)
            bp.space_before = Pt(2.5)
            bp.line_spacing = 1.15
            r = bp.add_run()
            r.text = "• " + bitem
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            r.font.color.rgb = TEXT_DARK

    # --- RIGHT: 3 Quantified Benefit Callout Cards ---
    benefits = [
        ("Clinical & Social Impact", "Compressing Door-to-Needle Intervals",
         "Dramatically reduces critical delays in acute stroke, STEMI, and trauma by eliminating secondary hospital transfers.\n(Saves ~1.9M brain neurons per minute saved in ischemic stroke).",
         Inches(1.25), GREEN_BG, GREEN_BORDER, GREEN_TEXT),

        ("Economic & Urban Efficiency", "Minimizing Fuel & Congestion Waste",
         "Dynamic shortest-feasible-path emergency routing eliminates circular transit, wasted vehicle idling, and citywide traffic disruption while optimizing ER throughput.",
         Inches(3.05), AMBER_BG, AMBER_BORDER, AMBER_TEXT),

        ("Forensic Governance", "Tamper-Evident SHA-256 Ledger",
         "Cryptographic policy hashing (hashPolicy) and immutable SHA-256 event audit trails provide forensic accountability for clinical governance and medico-legal auditability.",
         Inches(4.85), GRAY_BG, PRIMARY_BLUE, DARK_NAVY)
    ]

    for btitle, bsub, bdesc, t_pos, bg_c, bord_c, text_c in benefits:
        bcard = add_card(s5, Inches(8.50), t_pos, Inches(4.23), Inches(1.70), bg_color=bg_c, border_color=bord_c, border_width=1.5)
        btf = bcard.text_frame
        btf.margin_left = Inches(0.12)
        btf.margin_right = Inches(0.12)
        btf.margin_top = Inches(0.08)
        bp = btf.paragraphs[0]
        strip_bullet(bp)
        br1 = bp.add_run()
        br1.text = btitle + "\n"
        br1.font.name = "Arial"
        br1.font.size = Pt(11.5)
        br1.font.bold = True
        br1.font.color.rgb = text_c

        br2 = bp.add_run()
        br2.text = bsub + "\n"
        br2.font.name = "Arial"
        br2.font.size = Pt(9.5)
        br2.font.bold = True
        br2.font.color.rgb = PRIMARY_BLUE

        br3 = bp.add_run()
        br3.text = bdesc
        br3.font.name = "Arial"
        br3.font.size = Pt(8.5)
        br3.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 6: RESEARCH AND REFERENCES (2x2 CARDS & OPEN HUB)
    # =========================================================================
    s6 = prs.slides[5]
    for s in s6.shapes:
        if s.name == 'Title 1':
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "RESEARCH  AND REFERENCES"
            p.font.name = "Times New Roman"
            p.font.size = Pt(36)
            p.font.bold = True
            p.font.color.rgb = DARK_NAVY
        elif s.name == 'TextBox 8':
            s.left = Inches(0)
            s.top = Inches(0)
            s.width = Inches(0.1)
            s.height = Inches(0.1)
            s.text_frame.clear()

    # --- TOP: 4 Categorized Reference Cards (2x2 Grid) ---
    ref_cards = [
        ("Public Healthcare Registries", "Data Grounding & Facility Identity", [
            ("Karnataka State Health & Family Welfare Services: ", "https://hfw.karnataka.gov.in — official registry for state healthcare facilities."),
            ("BBMP Open Hospital Directory: ", "Bruhat Bengaluru Mahanagara Palike urban healthcare facility coordinates and publicly listed specialties."),
            ("7 Canonical Bengaluru Facilities: ", "Ground-truth benchmark dataset derived from open government registries; capacity marked UNKNOWN.")
        ], Inches(0.60), Inches(1.25)),

        ("Healthcare & Facility Standards", "Data Modeling & Nomenclature", [
            ("Registry of Hospitals in Network of Insurance (ROHINI): ", "Insurance Information Bureau of India (IIB) (https://iib.gov.in) — facility identification standards."),
            ("HL7 FHIR Release 4: ", "https://hl7.org/fhir — foundational standards for timestamped clinical resource representation and event provenance."),
            ("Emergency Classification: ", "Standardized care requirement taxonomies for stroke, cardiac arrest, trauma, and burn cases.")
        ], Inches(6.78), Inches(1.25)),

        ("Spatial Engine & Mapping Framework", "Geospatial Routing Architecture", [
            ("OpenStreetMap (OSM) Contributors: ", "https://www.openstreetmap.org — open-access topological road network vector graph."),
            ("MapLibre GL: ", "https://maplibre.org — hardware-accelerated WebGL client-side vector map renderer."),
            ("Valhalla & OSRM Routing Engines: ", "High-performance multimodal matrix calculation with safe fallback to UNKNOWN ETA.")
        ], Inches(0.60), Inches(3.40)),

        ("Cloud & Architectural Standards", "Event-Driven & AI Guardrails", [
            ("AWS Well-Architected Framework: ", "Healthcare Industry Lens & Event-Driven Architecture (Amazon EventBridge, Lambda, DynamoDB)."),
            ("DynamoDB Single-Table Design: ", "Low-latency sub-50ms operational state persistence architecture."),
            ("Advisory AI Guardrails: ", "Amazon Bedrock sandboxed to read-only clinical summaries; strictly excluded from routing logic.")
        ], Inches(6.78), Inches(3.40))
    ]

    for rtitle, rsub, rbullets, l_pos, t_pos in ref_cards:
        rcard = add_card(s6, l_pos, t_pos, Inches(5.95), Inches(2.05), bg_color=WHITE, border_color=BORDER_GRAY)
        rtf = rcard.text_frame
        rtf.margin_left = Inches(0.12)
        rtf.margin_right = Inches(0.12)
        rtf.margin_top = Inches(0.08)
        rp = rtf.paragraphs[0]
        strip_bullet(rp)
        rr1 = rp.add_run()
        rr1.text = rtitle + "\n"
        rr1.font.name = "Arial"
        rr1.font.size = Pt(11)
        rr1.font.bold = True
        rr1.font.color.rgb = DARK_NAVY

        rr2 = rp.add_run()
        rr2.text = rsub + "\n"
        rr2.font.name = "Arial"
        rr2.font.size = Pt(9)
        rr2.font.bold = True
        rr2.font.color.rgb = PRIMARY_BLUE

        for b_name, b_val in rbullets:
            bp = rtf.add_paragraph()
            strip_bullet(bp)
            bp.space_before = Pt(2)
            bp.line_spacing = 1.15
            r1 = bp.add_run()
            r1.text = "• " + b_name
            r1.font.name = "Arial"
            r1.font.size = Pt(8.5)
            r1.font.bold = True
            r1.font.color.rgb = TEXT_DARK
            r2 = bp.add_run()
            r2.text = b_val
            r2.font.name = "Arial"
            r2.font.size = Pt(8)
            r2.font.color.rgb = TEXT_MUTED

    # --- BOTTOM: Open Verification & Codebase Hub ---
    hub = add_card(s6, Inches(0.60), Inches(5.55), Inches(12.13), Inches(1.20), bg_color=WHITE, border_color=PRIMARY_BLUE, border_width=1.5)
    htf = hub.text_frame
    htf.margin_left = Inches(0.15)
    htf.margin_top = Inches(0.08)
    hp = htf.paragraphs[0]
    strip_bullet(hp)
    hr1 = hp.add_run()
    hr1.text = "PROJECT REPOSITORY & VERIFIED ARTIFACTS\n"
    hr1.font.name = "Arial"
    hr1.font.size = Pt(11.5)
    hr1.font.bold = True
    hr1.font.color.rgb = PRIMARY_BLUE

    hp2 = htf.add_paragraph()
    strip_bullet(hp2)
    hp2.space_before = Pt(2)
    r1 = hp2.add_run()
    r1.text = "• GitHub Repository: "
    r1.font.name = "Arial"
    r1.font.size = Pt(9.5)
    r1.font.bold = True
    r1.font.color.rgb = DARK_NAVY
    r2 = hp2.add_run()
    r2.text = "https://github.com/Zynx095/Jiva — Monorepo containing full API, 4 web apps, demo orchestrator, and CDK.\n"
    r2.font.name = "Arial"
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = TEXT_MUTED

    r3 = hp2.add_run()
    r3.text = "• Verification Baseline: "
    r3.font.name = "Arial"
    r3.font.size = Pt(9.5)
    r3.font.bold = True
    r3.font.color.rgb = DARK_NAVY
    r4 = hp2.add_run()
    r4.text = "25 automated test suites pass • 124 runtime checks verified • 7 Lambda bundles compiled • CDK synth passes."
    r4.font.name = "Arial"
    r4.font.size = Pt(9.5)
    r4.font.color.rgb = TEXT_MUTED

    # Save PPTX
    prs.save(out_pptx_path)
    print(f"Successfully generated visual PPTX at: {out_pptx_path}")

    # Export to PDF and PNG previews
    try:
        import win32com.client
        print("Exporting presentation to PDF and PNG previews via PowerPoint COM...")
        ppt_app = win32com.client.Dispatch("PowerPoint.Application")
        presentation = ppt_app.Presentations.Open(out_pptx_path, WithWindow=False)
        presentation.SaveAs(out_pdf_path, 32)
        print(f"PDF saved to: {out_pdf_path}")
        
        os.makedirs('ppt/slides_preview', exist_ok=True)
        for i in range(1, len(presentation.Slides) + 1):
            out_img = os.path.abspath(f'ppt/slides_preview/slide_{i}.png')
            presentation.Slides(i).Export(out_img, 'PNG', 1920, 1080)
        print("PNG previews updated successfully!")

        presentation.Close()
        ppt_app.Quit()
    except Exception as e:
        print(f"Notice: Export failed: {e}")

if __name__ == '__main__':
    build_presentation()
