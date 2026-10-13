import sys
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml import parse_xml

def strip_bullet(paragraph):
    """Explicitly removes any bullet formatting from a paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    for child in list(pPr):
        tag = child.tag.split('}')[-1]
        if tag.startswith('bu'):
            pPr.remove(child)
    pPr.append(parse_xml('<a:buNone xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'))

def build_presentation():
    src_path = os.path.abspath('ppt/IDE BOOTCAMP PPT FORMAT.pptx')
    out_pptx_path = os.path.abspath('ppt/JIVA_IDE_BOOTCAMP_FORMAT_FINAL.pptx')
    out_pdf_path = os.path.abspath('ppt/JIVA_IDE_BOOTCAMP_FORMAT_FINAL.pdf')

    print(f"Loading template from: {src_path}")
    prs = Presentation(src_path)

    # Colors matching the template
    PRIMARY_BLUE = RGBColor(0, 112, 192)    # #0070C0 from template bottom bar
    DARK_NAVY = RGBColor(0, 51, 102)        # Deep contrast navy for titles
    TEXT_DARK = RGBColor(31, 41, 55)         # Charcoal text
    TEXT_MUTED = RGBColor(75, 85, 99)       # Medium charcoal

    # =========================================================================
    # SLIDE 1: TITLE PAGE
    # =========================================================================
    s1 = prs.slides[0]
    for shape in s1.shapes:
        if shape.name == 'Title 7':
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "IDE BOOTCAMP 2026"
            p.font.name = "Garamond"
            p.font.size = Pt(40)
            p.font.bold = True
            p.font.color.rgb = DARK_NAVY
            
        elif shape.name == 'Subtitle 3':
            shape.left = Inches(1.00)
            shape.width = Inches(11.33)
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "JIVA: REAL-TIME HEALTHCARE COORDINATION"
            p.font.name = "Times New Roman"
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = PRIMARY_BLUE
            
        elif shape.name == 'TextBox 9':
            shape.left = Inches(0.36)
            shape.top = Inches(2.15)
            shape.width = Inches(5.75)
            shape.height = Inches(4.85)
            tf = shape.text_frame
            tf.word_wrap = True
            tf.clear()

            items = [
                ("• Problem Statement Title: ", "Real-Time Emergency Healthcare Coordination & Dynamic Hospital Routing"),
                ("• Theme: ", "Healthcare & MedTech / Cloud-Native Coordination Mesh"),
                ("• PS Category: ", "Software (Event-Driven Healthcare Mesh)"),
                ("• Team Name: ", "Monochrome"),
                ("• Team Members: ", "Yukith M Joseph  •  Pranathi Nagabhushana"),
                ("• Solution Core: ", "An event-driven layer maintaining a shared, timestamped care state for patients, ambulances, and hospitals, routing via explicit hospital acceptance and dynamically rerouting when state changes."),
                ("• Verification Status: ", "Fully functional local prototype verified on synthetic Bengaluru data (124/124 checks pass). AWS architecture modeled in CDK (not deployed).")
            ]

            for idx, (label, val) in enumerate(items):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                strip_bullet(p)
                p.space_before = Pt(5) if idx > 0 else Pt(0)
                p.space_after = Pt(2)
                p.line_spacing = 1.15
                
                r_lbl = p.add_run()
                r_lbl.text = label
                r_lbl.font.name = "Arial"
                r_lbl.font.size = Pt(13)
                r_lbl.font.bold = True
                r_lbl.font.color.rgb = PRIMARY_BLUE if idx < 5 else DARK_NAVY

                r_val = p.add_run()
                r_val.text = val
                r_val.font.name = "Arial"
                r_val.font.size = Pt(12) if idx >= 5 else Pt(13)
                r_val.font.bold = False if idx >= 5 else True
                r_val.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 2: IDEA TITLE
    # =========================================================================
    s2 = prs.slides[1]
    for shape in s2.shapes:
        if shape.name == 'Title 1':
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "JIVA: REAL-TIME HEALTHCARE COORDINATION"
            p.font.name = "Times New Roman"
            p.font.size = Pt(32)
            p.font.bold = True
            p.font.color.rgb = DARK_NAVY
            
        elif shape.name == 'TextBox 8':
            shape.left = Inches(0.67)
            shape.top = Inches(1.30)
            shape.width = Inches(12.00)
            shape.height = Inches(5.50)
            tf = shape.text_frame
            tf.word_wrap = True
            tf.clear()

            def add_header(text, is_first=False):
                p = tf.paragraphs[0] if is_first else tf.add_paragraph()
                strip_bullet(p)
                p.space_before = Pt(0) if is_first else Pt(8)
                p.space_after = Pt(2)
                r = p.add_run()
                r.text = text
                r.font.name = "Arial"
                r.font.size = Pt(14.5)
                r.font.bold = True
                r.font.color.rgb = PRIMARY_BLUE
                return p

            def add_bullet(prefix, text):
                p = tf.add_paragraph()
                strip_bullet(p)
                p.space_before = Pt(0)
                p.space_after = Pt(2)
                p.line_spacing = 1.15
                r1 = p.add_run()
                r1.text = "• " + prefix
                r1.font.name = "Arial"
                r1.font.size = Pt(12)
                r1.font.bold = True
                r1.font.color.rgb = TEXT_DARK
                r2 = p.add_run()
                r2.text = text
                r2.font.name = "Arial"
                r2.font.size = Pt(12)
                r2.font.bold = False
                r2.font.color.rgb = TEXT_MUTED
                return p

            add_header("Proposed Solution (Describe your Idea/Solution/Prototype)", is_first=True)
            add_bullet("Event-Driven Coordination Mesh: ", "JIVA coordinates emergency healthcare across patients, ambulances, hospitals, and dispatch coordinators without assuming static availability or intruding on hospital EHR systems.")

            add_header("Detailed explanation of the proposed solution")
            add_bullet("Four Role-Scoped Web Applications: ", "Dedicated interfaces for Patients, Paramedics, Hospital ED staff, and Regional Coordinators (+ Live Demo Orchestrator), ensuring tamper-proof, role-specific visibility.")
            add_bullet("Closed-Loop Hospital Acceptance Protocol: ", "Instead of assuming readiness from directory listings, JIVA dispatches a case-scoped acceptance request. Hospitals actively return ACCEPTED, LIMITED, REJECTED, or UNAVAILABLE.")
            add_bullet("Time-Limited Operational States: ", "Unanswered requests expire to UNKNOWN via deterministic timeouts, preventing ambulances from navigating toward unconfirmed emergency rooms.")

            add_header("How it addresses the problem")
            add_bullet("Eliminates Stale Information & Emergency Diversion: ", "Replaces fragmented phone calls and outdated bed counts with authenticated, timestamped event transitions and verifiable audit provenance.")
            add_bullet("Automated Dynamic Rerouting: ", "If an accepted hospital updates to UNAVAILABLE or a request expires en route, JIVA immediately triggers dynamic rerouting to the next best feasible candidate.")

            add_header("Innovation and uniqueness of the solution")
            add_bullet("No Presumed Availability: ", "A facility's listed ICU capability is never assumed to be available right now; explicit confirmation is mandatory.")
            add_bullet("Deterministic Shadow Feasibility: ", "Routing decisions rely on deterministic rule verification and policy hashing, isolating generative AI strictly to advisory, non-safety-critical case summaries.")

    # =========================================================================
    # SLIDE 3: TECHNICAL APPROACH
    # =========================================================================
    s3 = prs.slides[2]
    for shape in s3.shapes:
        if shape.name == 'Title 1':
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "TECHNICAL APPROACH"
            p.font.name = "Times New Roman"
            p.font.size = Pt(36)
            p.font.bold = True
            p.font.color.rgb = DARK_NAVY
            
        elif shape.name == 'TextBox 8':
            shape.left = Inches(0.67)
            shape.top = Inches(1.30)
            shape.width = Inches(12.00)
            shape.height = Inches(5.50)
            tf = shape.text_frame
            tf.word_wrap = True
            tf.clear()

            def add_header(text, is_first=False):
                p = tf.paragraphs[0] if is_first else tf.add_paragraph()
                strip_bullet(p)
                p.space_before = Pt(0) if is_first else Pt(8)
                p.space_after = Pt(2)
                r = p.add_run()
                r.text = text
                r.font.name = "Arial"
                r.font.size = Pt(14.5)
                r.font.bold = True
                r.font.color.rgb = PRIMARY_BLUE
                return p

            def add_bullet(prefix, text):
                p = tf.add_paragraph()
                strip_bullet(p)
                p.space_before = Pt(0)
                p.space_after = Pt(2)
                p.line_spacing = 1.15
                r1 = p.add_run()
                r1.text = "• " + prefix
                r1.font.name = "Arial"
                r1.font.size = Pt(12)
                r1.font.bold = True
                r1.font.color.rgb = TEXT_DARK
                r2 = p.add_run()
                r2.text = text
                r2.font.name = "Arial"
                r2.font.size = Pt(12)
                r2.font.bold = False
                r2.font.color.rgb = TEXT_MUTED
                return p

            add_header("Technologies to be used (e.g. programming languages, frameworks, hardware)", is_first=True)
            add_bullet("Full-Stack TypeScript & Node.js: ", "Backend event processing in Node.js with strict Zod runtime schema validation and bidirectional real-time state synchronization via Socket.IO.")
            add_bullet("Frontend Architecture: ", "React 19, Vite, and Tailwind CSS v4 powering 4 lightweight web apps with MapLibre GL for client-side spatial map rendering.")
            add_bullet("Spatial Routing & Transit Rules: ", "OpenStreetMap vector tiles with Valhalla/OSRM routing engines; route calculation failures safely fall back to UNKNOWN ETA (never ETA=0).")
            add_bullet("AWS Cloud Architecture (CDK-Modeled, Not Deployed): ", "Amazon EventBridge (healthcare event bus), AWS Lambda (event microservices), Amazon DynamoDB (operational state), Amazon API Gateway, Cognito, and Amazon Bedrock (advisory case summaries only).")
            add_bullet("Hardware & IoT Integration: ", "Physical vitals/ECG monitor hardware is currently on hold; simulated via standard ingestion payloads with zero mandatory hardware lock-in.")

            add_header("Methodology and process for implementation (Flow Charts/Images/ working prototype)")
            add_bullet("11-Stage Closed-Loop Coordination Pipeline: ", "[1] Healthcare Event → [2] Ingestion & Zod Validation → [3] Canonical Event Bus → [4] State Update & Audit Ledger → [5] Care Requirement Match → [6] Candidate Hospital Evaluation → [7] Case-Scoped Acceptance Request → [8] Hospital Response (ACCEPTED/LIMITED/REJECTED/UNAVAILABLE) → [9] Confirmed Destination Dispatch → [10] Real-Time State Broadcast (Socket.IO) → [11] Dynamic Reroute / Continuous Handover.")
            add_bullet("Fully Offline-Capable Local Prototype: ", "Self-contained local development server running offline on port-isolated Vite instances; clinical state held in memory for sub-50ms processing latency.")
            add_bullet("Automated Verification Suite: ", "25 unit suites passing, 124 runtime integration checks verified, and 7 AWS Lambda bundles successfully compiled with esbuild.")
            add_bullet("Shadow Feasibility Engine: ", "Deterministic rule engine running in parallel with 0 regression disagreements, enforcing evidence freshness and cryptographic policy hashing.")

    # =========================================================================
    # SLIDE 4: FEASIBILITY AND VIABILITY
    # =========================================================================
    s4 = prs.slides[3]
    for shape in s4.shapes:
        if shape.name == 'Title 1':
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "FEASIBILITY AND VIABILITY"
            p.font.name = "Times New Roman"
            p.font.size = Pt(36)
            p.font.bold = True
            p.font.color.rgb = DARK_NAVY
            
        elif shape.name == 'TextBox 8':
            shape.left = Inches(0.67)
            shape.top = Inches(1.30)
            shape.width = Inches(12.00)
            shape.height = Inches(5.50)
            tf = shape.text_frame
            tf.word_wrap = True
            tf.clear()

            def add_header(text, is_first=False):
                p = tf.paragraphs[0] if is_first else tf.add_paragraph()
                strip_bullet(p)
                p.space_before = Pt(0) if is_first else Pt(8)
                p.space_after = Pt(2)
                r = p.add_run()
                r.text = text
                r.font.name = "Arial"
                r.font.size = Pt(14.5)
                r.font.bold = True
                r.font.color.rgb = PRIMARY_BLUE
                return p

            def add_bullet(prefix, text):
                p = tf.add_paragraph()
                strip_bullet(p)
                p.space_before = Pt(0)
                p.space_after = Pt(2)
                p.line_spacing = 1.15
                r1 = p.add_run()
                r1.text = "• " + prefix
                r1.font.name = "Arial"
                r1.font.size = Pt(12)
                r1.font.bold = True
                r1.font.color.rgb = TEXT_DARK
                r2 = p.add_run()
                r2.text = text
                r2.font.name = "Arial"
                r2.font.size = Pt(12)
                r2.font.bold = False
                r2.font.color.rgb = TEXT_MUTED
                return p

            add_header("Analysis of the feasibility of the idea", is_first=True)
            add_bullet("Lightweight Overlay Model: ", "Requires zero invasive modifications to hospital electronic health records (EHR); operates as an external, secure coordination layer accessible via standard web browsers.")
            add_bullet("Demonstrated Execution: ", "Local architecture processes simultaneous multi-party events with sub-50ms latency across 7 canonical Bengaluru hospital facilities.")
            add_bullet("Scalable Cloud Pathway: ", "CDK-synthesized architecture maps each local component 1:1 to AWS serverless primitives (EventBridge, Lambda, DynamoDB), providing predictable, cost-effective pay-per-event scaling.")

            add_header("Potential challenges and risks")
            add_bullet("Public Directory Blind Spots: ", "Karnataka / BBMP public hospital directories list registered facilities but lack real-time bed or specialist status.")
            add_bullet("Transit Connectivity Loss: ", "Ambulances frequently experience intermittent cellular drops while navigating urban transit corridors.")
            add_bullet("In-Memory State Volatility: ", "Local prototype relies on in-memory storage; process termination resets active emergency session data.")
            add_bullet("Clinical AI Safety: ", "LLM hallucinations pose severe safety hazards if used for dispatch, routing, or triage calculations.")

            add_header("Strategies for overcoming these challenges")
            add_bullet("Honest Availability Semantics: ", "Directory listings without verified live telemetry are flagged UNKNOWN. Acceptance expires to UNKNOWN rather than defaulting to available.")
            add_bullet("Offline-Resilient Clients: ", "Client-side state caching with automatic Socket.IO reconnection backoff; navigation preserves last known location with explicit UNKNOWN ETA.")
            add_bullet("Persistence Architecture: ", "Structured migration roadmap to transactional SQLite for robust local edge storage, and DynamoDB single-table design for AWS deployment.")
            add_bullet("Strict AI Decision Bounding: ", "Amazon Bedrock is strictly sandboxed to read-only clinical handoff summaries; routing and dispatch remain 100% deterministic code.")

    # =========================================================================
    # SLIDE 5: IMPACT AND BENEFITS
    # =========================================================================
    s5 = prs.slides[4]
    for shape in s5.shapes:
        if shape.name == 'Title 1':
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "IMPACT AND BENEFITS"
            p.font.name = "Times New Roman"
            p.font.size = Pt(36)
            p.font.bold = True
            p.font.color.rgb = DARK_NAVY
            
        elif shape.name == 'TextBox 8':
            shape.left = Inches(0.67)
            shape.top = Inches(1.30)
            shape.width = Inches(12.00)
            shape.height = Inches(5.50)
            tf = shape.text_frame
            tf.word_wrap = True
            tf.clear()

            def add_header(text, is_first=False):
                p = tf.paragraphs[0] if is_first else tf.add_paragraph()
                strip_bullet(p)
                p.space_before = Pt(0) if is_first else Pt(8)
                p.space_after = Pt(2)
                r = p.add_run()
                r.text = text
                r.font.name = "Arial"
                r.font.size = Pt(14.5)
                r.font.bold = True
                r.font.color.rgb = PRIMARY_BLUE
                return p

            def add_bullet(prefix, text):
                p = tf.add_paragraph()
                strip_bullet(p)
                p.space_before = Pt(0)
                p.space_after = Pt(2)
                p.line_spacing = 1.15
                r1 = p.add_run()
                r1.text = "• " + prefix
                r1.font.name = "Arial"
                r1.font.size = Pt(12)
                r1.font.bold = True
                r1.font.color.rgb = TEXT_DARK
                r2 = p.add_run()
                r2.text = text
                r2.font.name = "Arial"
                r2.font.size = Pt(12)
                r2.font.bold = False
                r2.font.color.rgb = TEXT_MUTED
                return p

            add_header("Potential impact on the target audience", is_first=True)
            add_bullet("Emergency Patients & Families: ", "Eliminates dangerous ER turnaways and secondary transfers; provides transparent status tracking and honest travel ETAs during critical care intervals.")
            add_bullet("Ambulance Crews & Paramedics: ", "Removes the chaos of manual en-route phone calls, guaranteeing a confirmed hospital destination with staff prepped before arrival.")
            add_bullet("Hospital Emergency Departments: ", "Provides structured pre-arrival clinical telemetry, enabling EDs to prepare trauma bays, cath labs, or surgical teams in advance.")
            add_bullet("Regional Emergency Coordinators: ", "Delivers a live, centralized operational view of hospital bed states, pending acceptance requests, and regional ambulance dispatches.")

            add_header("Benefits of the solution (social, economic, environmental, etc.)")
            add_bullet("Compressing Door-to-Needle Intervals: ", "Dramatically reduces critical delays for stroke, STEMI, and polytrauma cases by eliminating hospital shopping and redirecting ambulances dynamically.")
            add_bullet("Economic Hospital Efficiency: ", "Prevents costly emergency room congestion, uncoordinated ambulance diversions, and wasted ICU prep by balancing regional emergency intake.")
            add_bullet("Environmental & Fleet Optimization: ", "Reduces wasted ambulance transit mileage, vehicle idling, and city traffic disruption through dynamic shortest-feasible-path routing.")
            add_bullet("Immutable Forensic Accountability: ", "Cryptographic policy hashing (hashPolicy) and SHA-256 event audit logging produce an unalterable forensic record for clinical and medico-legal auditability.")

    # =========================================================================
    # SLIDE 6: RESEARCH  AND REFERENCES
    # =========================================================================
    s6 = prs.slides[5]
    for shape in s6.shapes:
        if shape.name == 'Title 1':
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "RESEARCH  AND REFERENCES"
            p.font.name = "Times New Roman"
            p.font.size = Pt(36)
            p.font.bold = True
            p.font.color.rgb = DARK_NAVY
            
        elif shape.name == 'TextBox 8':
            shape.left = Inches(0.67)
            shape.top = Inches(1.30)
            shape.width = Inches(12.00)
            shape.height = Inches(5.50)
            tf = shape.text_frame
            tf.word_wrap = True
            tf.clear()

            def add_header(text, is_first=False):
                p = tf.paragraphs[0] if is_first else tf.add_paragraph()
                strip_bullet(p)
                p.space_before = Pt(0) if is_first else Pt(8)
                p.space_after = Pt(2)
                r = p.add_run()
                r.text = text
                r.font.name = "Arial"
                r.font.size = Pt(14.5)
                r.font.bold = True
                r.font.color.rgb = PRIMARY_BLUE
                return p

            def add_bullet(prefix, text):
                p = tf.add_paragraph()
                strip_bullet(p)
                p.space_before = Pt(0)
                p.space_after = Pt(2)
                p.line_spacing = 1.15
                r1 = p.add_run()
                r1.text = "• " + prefix
                r1.font.name = "Arial"
                r1.font.size = Pt(12)
                r1.font.bold = True
                r1.font.color.rgb = TEXT_DARK
                r2 = p.add_run()
                r2.text = text
                r2.font.name = "Arial"
                r2.font.size = Pt(12)
                r2.font.bold = False
                r2.font.color.rgb = TEXT_MUTED
                return p

            add_header("Details / Links of the reference and research work", is_first=True)
            add_bullet("Government Healthcare Registries: ", "Karnataka State Health & Family Welfare Services (https://hfw.karnataka.gov.in) and BBMP Open Hospital Directory — source registries for 7 canonical Bengaluru hospital facilities and publicly listed emergency capabilities.")
            add_bullet("Insurance & Health Facility Standards: ", "Registry of Hospitals in Network of Insurance (ROHINI), Insurance Information Bureau of India (IIB) (https://iib.gov.in) — standard identifiers and tier classification schemas.")
            add_bullet("Open Spatial Mapping & Routing: ", "OpenStreetMap contributors (https://www.openstreetmap.org) for geospatial road networks; MapLibre GL (https://maplibre.org) for client-side WebGL rendering; Valhalla & OSRM routing engine architectures.")
            add_bullet("Event-Driven Architecture Standards: ", "AWS Well-Architected Framework: Healthcare Industry Lens (Amazon EventBridge event patterns, AWS Lambda concurrency, DynamoDB single-table design).")
            add_bullet("Healthcare Interoperability Paradigms: ", "HL7 FHIR (Fast Healthcare Interoperability Resources) Release 4 (https://hl7.org/fhir) — foundational principles for timestamped clinical resource representation and provenance tracking.")
            add_bullet("JIVA Open Source Repository: ", "https://github.com/Zynx095/Jiva — complete source code, 25 test suites, 124 runtime checks, CDK infrastructure definitions, and reproducibility documentation.")

    # Save final PPTX
    prs.save(out_pptx_path)
    print(f"Successfully generated final presentation at: {out_pptx_path}")

    # Convert to PDF and export PNGs
    try:
        import win32com.client
        print("Exporting presentation to PDF and PNGs via PowerPoint COM...")
        ppt_app = win32com.client.Dispatch("PowerPoint.Application")
        presentation = ppt_app.Presentations.Open(out_pptx_path, WithWindow=False)
        presentation.SaveAs(out_pdf_path, 32)
        print(f"PDF saved to: {out_pdf_path}")
        
        os.makedirs('ppt/slides_preview', exist_ok=True)
        for i in range(1, len(presentation.Slides) + 1):
            out_img = os.path.abspath(f'ppt/slides_preview/slide_{i}.png')
            presentation.Slides(i).Export(out_img, 'PNG', 1920, 1080)
        print("PNG previews updated successfully!")

        presentation.Close()
        ppt_app.Quit()
    except Exception as e:
        print(f"Notice: Export failed: {e}")

if __name__ == '__main__':
    build_presentation()
