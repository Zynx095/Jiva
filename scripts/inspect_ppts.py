import pptx
from pptx import Presentation

def inspect_prs(prs_path, name, out_file):
    out_file.write(f"\n=======================================================\n")
    out_file.write(f"INSPECTING: {name} ({prs_path})\n")
    out_file.write(f"=======================================================\n")
    prs = Presentation(prs_path)
    out_file.write(f"Slide Dimensions: {prs.slide_width / 914400:.3f} x {prs.slide_height / 914400:.3f} inches\n")
    out_file.write(f"Slide Count: {len(prs.slides)}\n")
    
    for idx, slide in enumerate(prs.slides):
        out_file.write(f"\n--- SLIDE {idx+1} [Layout: {slide.slide_layout.name}] ---\n")
        # Check background
        try:
            bg = slide.background
            out_file.write(f"  Background: {bg}\n")
        except:
            pass
        for s_idx, shape in enumerate(slide.shapes):
            stype = str(shape.shape_type)
            pos = f"L={shape.left/914400:.2f}, T={shape.top/914400:.2f}, W={shape.width/914400:.2f}, H={shape.height/914400:.2f}"
            out_file.write(f"  Shape {s_idx} [{shape.name} | {stype} | {pos}]\n")
            
            if shape.has_text_frame:
                tf = shape.text_frame
                for p_idx, p in enumerate(tf.paragraphs):
                    if not p.text.strip():
                        continue
                    runs_details = []
                    for r in p.runs:
                        font_name = r.font.name or "inherited"
                        font_size = f"{r.font.size.pt}pt" if r.font.size else "inherited"
                        bold = "BOLD" if r.font.bold else ""
                        color = ""
                        try:
                            if r.font.color and r.font.color.rgb:
                                color = f"rgb({r.font.color.rgb})"
                        except:
                            pass
                        runs_details.append(f"'{r.text}' ({font_name}, {font_size}, {bold}, {color})")
                    out_file.write(f"    P{p_idx} [level={p.level}, align={p.alignment}]: {p.text}\n")
                    if runs_details:
                        out_file.write(f"       Runs: {runs_details}\n")
                        
            if shape.has_table:
                table = shape.table
                out_file.write(f"    Table: {len(table.rows)} rows x {len(table.columns)} cols\n")
                for r_idx, row in enumerate(table.rows):
                    row_vals = [c.text.replace('\n', ' ') for c in row.cells]
                    out_file.write(f"      Row {r_idx}: {row_vals}\n")

with open('ppt_inspection_utf8.txt', 'w', encoding='utf-8') as f:
    inspect_prs('ppt/IDE BOOTCAMP PPT FORMAT.pptx', 'IDE BOOTCAMP PPT FORMAT', f)
    inspect_prs('ppt/JIVA-AWS-Innovation-Challenge-2026-DRAFT.pptx', 'JIVA-AWS-Innovation-Challenge-2026-DRAFT', f)
    inspect_prs('ppt/AWS-INNOVATION-CHALLENGE-2026.pptx', 'AWS-INNOVATION-CHALLENGE-2026', f)

print("Done writing ppt_inspection_utf8.txt")
