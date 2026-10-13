from pptx import Presentation

prs = Presentation('ppt/IDE BOOTCAMP PPT FORMAT.pptx')
with open('ide_format_runs.txt', 'w', encoding='utf-8') as f:
    for idx, slide in enumerate(prs.slides):
        f.write(f"=== SLIDE {idx+1} ===\n")
        for s in slide.shapes:
            if s.has_text_frame:
                f.write(f"  Shape {s.name} (L={s.left/914400:.2f}, T={s.top/914400:.2f}, W={s.width/914400:.2f}, H={s.height/914400:.2f}):\n")
                for p_idx, p in enumerate(s.text_frame.paragraphs):
                    f.write(f"    P{p_idx} (space_before={p.space_before}, space_after={p.space_after}, line_spacing={p.line_spacing}, level={p.level}, align={p.alignment}):\n")
                    for r in p.runs:
                        fn = r.font.name
                        fs = r.font.size.pt if r.font.size else None
                        b = r.font.bold
                        c = str(r.font.color.rgb) if r.font.color and hasattr(r.font.color, 'rgb') and r.font.color.rgb else 'theme/none'
                        f.write(f"      run: '{r.text}' | font={fn}, size={fs}, bold={b}, color={c}\n")

print("Done writing ide_format_runs.txt")
