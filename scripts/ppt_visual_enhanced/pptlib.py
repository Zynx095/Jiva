"""Editable PowerPoint drawing helpers and a small vector illustration set for the JIVA deck.

Everything is drawn with native PowerPoint shapes (autoshapes, freeforms, connectors, text boxes) so the
result stays editable. Colours and fonts follow the existing IDE Bootcamp deck.
"""
import math

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

# ---------------------------------------------------------------- palette (taken from the deck)
NAVY = '003366'
BLUE = '0070C0'
TINT = 'E3EEF9'        # light tint of the template blue (screens, windows)
SLATE_BG = 'F8FAFC'
SLATE = 'CBD5E1'
TEXT = '1F2937'
MUTED = '4B5563'
WHITE = 'FFFFFF'
GREEN, GREEN_BG, GREEN_TX = '16A34A', 'DCFCE7', '15803D'
AMBER, AMBER_BG, AMBER_TX = 'D97706', 'FEF3C7', 'B45309'
AMBER_TINT = 'FFFBEB'
RED, RED_BG, RED_TX = 'DC2626', 'FEE2E2', 'B91C1C'
GREY, GREY_BG, GREY_TX = '9CA3AF', 'F3F4F6', '4B5563'
FILM = '0F2A47'        # dark film for imaging thumbnails (a deeper navy)
FILM_LIGHT = 'D6E2EF'  # bone / echo highlights on film
GRID = 'E2E8F0'

DASH = MSO_LINE_DASH_STYLE.DASH
SQUARE_DOT = MSO_LINE_DASH_STYLE.SQUARE_DOT

EMU_IN = 914400


def E(v):
    return int(round(v * EMU_IN))


def rgb(s):
    return RGBColor.from_string(s)


# ---------------------------------------------------------------- low-level shape helpers
_SPPR_BEFORE_EFFECT = {qn(t) for t in ('a:xfrm', 'a:custGeom', 'a:prstGeom', 'a:noFill', 'a:solidFill', 'a:gradFill',
                                        'a:blipFill', 'a:pattFill', 'a:grpFill', 'a:ln')}


def no_effects(shape):
    """Override the theme's shape effect (soft shadow) with an explicit empty effect list."""
    spPr = shape._element.spPr
    for e in spPr.findall(qn('a:effectLst')):
        spPr.remove(e)
    idx = 0
    for i, child in enumerate(spPr):
        if child.tag in _SPPR_BEFORE_EFFECT:
            idx = i + 1
    spPr.insert(idx, etree.Element(qn('a:effectLst')))


def _style_line(line, color, lw, dash):
    if color is None:
        line.fill.background()
    else:
        line.color.rgb = rgb(color)
        line.width = Pt(lw)
        if dash is not None:
            line.dash_style = dash


def _ends(line, head=None, tail=None, size='med'):
    ln = line._get_or_add_ln()
    for tag in ('a:headEnd', 'a:tailEnd'):
        for e in ln.findall(qn(tag)):
            ln.remove(e)
    if head:
        etree.SubElement(ln, qn('a:headEnd'), type=head, w=size, len=size)
    if tail:
        etree.SubElement(ln, qn('a:tailEnd'), type=tail, w=size, len=size)


def shp(shapes, kind, x, y, w, h, fill=None, line=None, lw=1.0, dash=None, shadow=False, adj=None, rot=0.0, name=None):
    s = shapes.add_shape(kind, E(x), E(y), E(w), E(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    _style_line(s.line, line, lw, dash)
    if adj:
        for i, v in enumerate(adj):
            if v is not None:
                s.adjustments[i] = v
    if rot:
        s.rotation = rot
    if not shadow:
        no_effects(s)
    if name:
        s.name = name
    return s


def poly(shapes, pts, line=NAVY, lw=1.0, fill=None, closed=False, dash=None, head=None, tail=None, name=None, arrow_size='med'):
    """Freeform polyline/polygon through absolute points (inches)."""
    fb = shapes.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
    fb.add_line_segments([(E(px), E(py)) for px, py in pts[1:]], close=closed)
    s = fb.convert_to_shape()
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    _style_line(s.line, line, lw, dash)
    if head or tail:
        _ends(s.line, head, tail, arrow_size)
    no_effects(s)
    if name:
        s.name = name
    return s


def arrow(shapes, x1, y1, x2, y2, color=BLUE, lw=1.5, dash=None, head=None, tail='triangle', name=None, size='med'):
    """Straight connector with an arrowhead at the end (x2, y2)."""
    c = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    _style_line(c.line, color, lw, dash)
    _ends(c.line, head, tail, size)
    if name:
        c.name = name
    return c


def line(shapes, x1, y1, x2, y2, color=NAVY, lw=1.0, dash=None, name=None):
    return arrow(shapes, x1, y1, x2, y2, color=color, lw=lw, dash=dash, tail=None, name=name)


def group(shapes, name=None):
    g = shapes.add_group_shape()
    if name:
        g.name = name
    return g


ALIGN = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}
ANCHOR = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}


def fill_text(tf, paras, size=10, color=TEXT, bold=False, italic=False, align='c', anchor='m', font='Arial',
              ins=(0.04, 0.03, 0.04, 0.03), wrap=True, space_before=0, line_spacing=None):
    """paras: list of paragraphs; each is a str or a list of runs (text, {size,bold,italic,color,font})."""
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [E(v) for v in ins]
    tf.vertical_anchor = ANCHOR[anchor]
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        popts = {}
        if isinstance(para, tuple) and len(para) == 2 and isinstance(para[1], dict) and isinstance(para[0], list):
            para, popts = para
        p.alignment = ALIGN[popts.get('align', align)]
        sb = popts.get('space_before', 0 if first else space_before)
        if sb:
            p.space_before = Pt(sb)
        ls = popts.get('line_spacing', line_spacing)
        if ls:
            p.line_spacing = ls
        runs = [(para, {})] if isinstance(para, str) else para
        for text, o in runs:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = o.get('font', font)
            f.size = Pt(o.get('size', size))
            f.bold = o.get('bold', bold)
            f.italic = o.get('italic', italic)
            f.color.rgb = rgb(o.get('color', color))
        first = False
    return tf


def text(shapes, x, y, w, h, paras, name=None, **kw):
    tb = shapes.add_textbox(E(x), E(y), E(w), E(h))
    kw.setdefault('ins', (0.0, 0.0, 0.0, 0.0))
    fill_text(tb.text_frame, paras, **kw)
    if name:
        tb.name = name
    return tb


def card(shapes, x, y, w, h, fill=WHITE, line=SLATE, lw=1.0, radius=0.12, shadow=True, dash=None, name=None,
         paras=None, **kw):
    """Rounded card in the deck's style. `radius` is the corner radius in inches."""
    s = shp(shapes, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line, lw=lw, dash=dash, shadow=shadow,
            name=name)
    s.adjustments[0] = min(0.5, radius / min(w, h))
    if paras is not None:
        fill_text(s.text_frame, paras, **kw)
    return s


def pill(shapes, x, y, w, h, label, fill, border, color, size=8, bold=True, name=None, lw=0.75, ins=0.03):
    s = shp(shapes, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=border, lw=lw, name=name)
    s.adjustments[0] = 0.5
    fill_text(s.text_frame, [label], size=size, color=color, bold=bold, ins=(ins, 0.0, ins, 0.0), wrap=False)
    return s


STATUS = {
    'done': (GREEN_BG, GREEN, GREEN_TX),
    'dev': (AMBER_BG, AMBER, AMBER_TX),
    'cdk': ('E8F1FA', BLUE, BLUE),
    'plan': (GREY_BG, GREY, GREY_TX),
    'neutral': (WHITE, SLATE, MUTED),
}


def status_pill(shapes, x, y, w, h, label, kind, size=8, name=None):
    f, b, c = STATUS[kind]
    return pill(shapes, x, y, w, h, label, f, b, c, size=size, name=name)


def num_badge(shapes, cx, cy, d, n, fill=BLUE, color=WHITE, size=11):
    s = shp(shapes, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, fill=fill, line=WHITE, lw=1.25, name=f'Step {n} number')
    fill_text(s.text_frame, [str(n)], size=size, color=color, bold=True, ins=(0, 0, 0, 0))
    return s


# ---------------------------------------------------------------- illustration set
def _lw(h, lo=0.75, hi=1.5, k=2.4):
    return max(lo, min(hi, h * k))


def badge(shapes, cx, cy, d, kind, name=None):
    """Status badge: ok (green check), no (red cross), busy (amber bar), expired (grey clock), wait (grey dots)."""
    fills = {'ok': GREEN, 'no': RED, 'busy': AMBER, 'expired': GREY, 'wait': GREY}
    g = group(shapes, name or f'Badge {kind}')
    gs = g.shapes
    x, y = cx - d / 2, cy - d / 2
    shp(gs, MSO_SHAPE.OVAL, x, y, d, d, fill=fills[kind], line=WHITE, lw=1.0)
    w = max(1.0, d * 9.5)
    P = lambda fx, fy: (x + fx * d, y + fy * d)
    if kind == 'ok':
        poly(gs, [P(0.26, 0.53), P(0.43, 0.70), P(0.75, 0.33)], line=WHITE, lw=w)
    elif kind == 'no':
        line(gs, *P(0.32, 0.32), *P(0.68, 0.68), color=WHITE, lw=w)
        line(gs, *P(0.68, 0.32), *P(0.32, 0.68), color=WHITE, lw=w)
    elif kind == 'busy':
        shp(gs, MSO_SHAPE.RECTANGLE, x + 0.24 * d, y + 0.43 * d, 0.52 * d, 0.14 * d, fill=WHITE)
    elif kind == 'expired':
        shp(gs, MSO_SHAPE.OVAL, x + 0.2 * d, y + 0.2 * d, 0.6 * d, 0.6 * d, fill=None, line=WHITE, lw=max(0.75, w * 0.6))
        line(gs, *P(0.5, 0.5), *P(0.5, 0.31), color=WHITE, lw=max(0.75, w * 0.6))
        line(gs, *P(0.5, 0.5), *P(0.64, 0.5), color=WHITE, lw=max(0.75, w * 0.6))
    elif kind == 'wait':
        for fx in (0.3, 0.5, 0.7):
            shp(gs, MSO_SHAPE.OVAL, x + (fx - 0.06) * d, y + 0.44 * d, 0.12 * d, 0.12 * d, fill=WHITE)
    return g


def ambulance(shapes, x, y, h, name='Ambulance illustration'):
    W = h * 1.75
    g = group(shapes, name)
    gs = g.shapes
    P = lambda fx, fy: (x + fx * W, y + fy * h)
    lw = _lw(h)
    shp(gs, MSO_SHAPE.ROUNDED_RECTANGLE, *P(0.29, 0.03), 0.13 * W, 0.11 * h, fill=BLUE, adj=[0.3])
    body = shp(gs, MSO_SHAPE.ROUNDED_RECTANGLE, *P(0.0, 0.13), 0.645 * W, 0.67 * h, fill=WHITE, line=NAVY, lw=lw)
    body.adjustments[0] = 0.10
    poly(gs, [P(0.62, 0.31), P(0.80, 0.31), P(0.955, 0.52), P(1.0, 0.56), P(1.0, 0.80), P(0.62, 0.80)], line=NAVY, lw=lw,
         fill=WHITE, closed=True)
    poly(gs, [P(0.68, 0.375), P(0.785, 0.375), P(0.895, 0.53), P(0.68, 0.53)], line=None, fill=TINT, closed=True)
    shp(gs, MSO_SHAPE.RECTANGLE, *P(0.006, 0.60), 0.988 * W, 0.065 * h, fill=BLUE)
    c = 0.30 * h
    cx, cy = P(0.32, 0.37)
    shp(gs, MSO_SHAPE.CROSS, cx - c / 2, cy - c / 2, c, c, fill=BLUE, adj=[0.3])
    for fx in (0.20, 0.79):
        wx, wy = P(fx, 0.82)
        d = 0.29 * h
        shp(gs, MSO_SHAPE.OVAL, wx - d / 2, wy - d / 2, d, d, fill=NAVY, line=WHITE, lw=0.75)
        d2 = 0.11 * h
        shp(gs, MSO_SHAPE.OVAL, wx - d2 / 2, wy - d2 / 2, d2, d2, fill=WHITE)
    return g


def ambulance_width(h):
    return h * 1.75


def hospital(shapes, x, y, h, status=None, name='Hospital illustration', muted=False):
    W = h * 1.15
    g = group(shapes, name)
    gs = g.shapes
    P = lambda fx, fy: (x + fx * W, y + fy * h)
    lw = _lw(h)
    edge = GREY if muted else NAVY
    accent = GREY if muted else BLUE
    win = GREY_BG if muted else TINT
    shp(gs, MSO_SHAPE.RECTANGLE, *P(0.02, 0.36), 0.34 * W, 0.64 * h, fill=WHITE, line=edge, lw=lw)
    shp(gs, MSO_SHAPE.RECTANGLE, *P(0.64, 0.36), 0.34 * W, 0.64 * h, fill=WHITE, line=edge, lw=lw)
    shp(gs, MSO_SHAPE.RECTANGLE, *P(0.30, 0.08), 0.40 * W, 0.92 * h, fill=WHITE, line=edge, lw=lw)
    c = 0.25 * h
    cx, cy = P(0.50, 0.27)
    shp(gs, MSO_SHAPE.CROSS, cx - c / 2, cy - c / 2, c, c, fill=accent, adj=[0.3])
    for fx in (0.09, 0.71):
        for fy in (0.46, 0.66):
            shp(gs, MSO_SHAPE.RECTANGLE, *P(fx, fy), 0.20 * W, 0.11 * h, fill=win)
    shp(gs, MSO_SHAPE.RECTANGLE, *P(0.38, 0.50), 0.24 * W, 0.10 * h, fill=win)
    shp(gs, MSO_SHAPE.RECTANGLE, *P(0.415, 0.74), 0.17 * W, 0.26 * h, fill=accent)
    if status:
        d = 0.40 * h
        bx, by = P(0.93, 0.13)
        badge(gs, bx, by, d, status)
    return g


def hospital_width(h):
    return h * 1.15


def _shoulders(gs, P, fill, lw, edge=NAVY):
    return poly(gs, [P(0.08, 1.0), P(0.10, 0.80), P(0.16, 0.665), P(0.26, 0.575), P(0.40, 0.525), P(0.60, 0.525),
                     P(0.74, 0.575), P(0.84, 0.665), P(0.90, 0.80), P(0.92, 1.0)], line=edge, lw=lw, fill=fill,
                closed=True)


def doctor(shapes, x, y, h, name='Doctor illustration'):
    W = h
    g = group(shapes, name)
    gs = g.shapes
    P = lambda fx, fy: (x + fx * W, y + fy * h)
    lw = _lw(h)
    _shoulders(gs, P, WHITE, lw)
    poly(gs, [P(0.405, 0.53), P(0.595, 0.53), P(0.50, 0.71)], line=None, fill=BLUE, closed=True)
    line(gs, *P(0.40, 0.53), *P(0.47, 0.86), color=NAVY, lw=max(0.5, lw * 0.6))
    line(gs, *P(0.60, 0.53), *P(0.53, 0.86), color=NAVY, lw=max(0.5, lw * 0.6))
    poly(gs, [P(0.37, 0.55), P(0.30, 0.64), P(0.285, 0.74), P(0.31, 0.80)], line=NAVY, lw=lw)
    d = 0.11 * h
    sx, sy = P(0.315, 0.83)
    shp(gs, MSO_SHAPE.OVAL, sx - d / 2, sy - d / 2, d, d, fill=BLUE, line=NAVY, lw=max(0.5, lw * 0.6))
    c = 0.11 * h
    cx, cy = P(0.71, 0.80)
    shp(gs, MSO_SHAPE.CROSS, cx - c / 2, cy - c / 2, c, c, fill=BLUE, adj=[0.3])
    shp(gs, MSO_SHAPE.OVAL, *P(0.32, 0.05), 0.36 * W, 0.41 * h, fill=WHITE, line=NAVY, lw=lw)
    return g


def patient(shapes, x, y, h, name='Patient illustration'):
    W = h
    g = group(shapes, name)
    gs = g.shapes
    P = lambda fx, fy: (x + fx * W, y + fy * h)
    lw = _lw(h)
    _shoulders(gs, P, BLUE, lw)
    hs = 0.22 * h
    hx, hy = P(0.50, 0.77)
    shp(gs, MSO_SHAPE.HEART, hx - hs / 2, hy - hs / 2, hs, hs * 0.9, fill=WHITE)
    shp(gs, MSO_SHAPE.OVAL, *P(0.32, 0.05), 0.36 * W, 0.41 * h, fill=WHITE, line=NAVY, lw=lw)
    return g


def person(shapes, x, y, h, fill=NAVY, name='Person illustration'):
    """Plain person silhouette (crew member, coordinator)."""
    W = h
    g = group(shapes, name)
    gs = g.shapes
    P = lambda fx, fy: (x + fx * W, y + fy * h)
    lw = _lw(h)
    _shoulders(gs, P, fill, lw, edge=fill)
    shp(gs, MSO_SHAPE.OVAL, *P(0.32, 0.05), 0.36 * W, 0.41 * h, fill=fill, line=fill, lw=lw)
    return g


def monitor(shapes, x, y, h, content='dashboard', name='Dashboard illustration'):
    W = h * 1.25
    g = group(shapes, name)
    gs = g.shapes
    P = lambda fx, fy: (x + fx * W, y + fy * h)
    lw = _lw(h)
    scr = shp(gs, MSO_SHAPE.ROUNDED_RECTANGLE, *P(0.02, 0.02), 0.96 * W, 0.72 * h, fill=WHITE, line=NAVY, lw=lw)
    scr.adjustments[0] = 0.07
    shp(gs, MSO_SHAPE.RECTANGLE, *P(0.08, 0.10), 0.84 * W, 0.56 * h, fill=TINT)
    if content == 'dashboard':
        base = 0.60
        for fx, top, col in ((0.14, 0.42, NAVY), (0.26, 0.30, BLUE), (0.38, 0.20, NAVY)):
            shp(gs, MSO_SHAPE.RECTANGLE, *P(fx, top), 0.08 * W, (base - top) * h, fill=col)
        pd = 0.20 * h
        px, py = P(0.71, 0.17)
        shp(gs, MSO_SHAPE.OVAL, px - pd / 2, py, pd, pd, fill=BLUE)
        poly(gs, [(px - pd * 0.42, py + pd * 0.72), (px + pd * 0.42, py + pd * 0.72), (px, py + pd * 1.45)], line=None,
             fill=BLUE, closed=True)
        dd = pd * 0.38
        shp(gs, MSO_SHAPE.OVAL, px - dd / 2, py + pd / 2 - dd / 2, dd, dd, fill=WHITE)
    elif content == 'play':
        t = 0.30 * h
        tx, ty = P(0.36, 0.23)
        shp(gs, MSO_SHAPE.ISOSCELES_TRIANGLE, tx, ty, t, t, fill=AMBER, rot=90)
        for fy in (0.30, 0.42, 0.54):
            shp(gs, MSO_SHAPE.RECTANGLE, *P(0.62, fy), 0.20 * W, 0.045 * h, fill=GREY)
    elif content == 'lines':
        for fy, fw in ((0.20, 0.6), (0.32, 0.7), (0.44, 0.5)):
            shp(gs, MSO_SHAPE.RECTANGLE, *P(0.16, fy), fw * W, 0.06 * h, fill=SLATE)
    shp(gs, MSO_SHAPE.RECTANGLE, *P(0.45, 0.74), 0.10 * W, 0.12 * h, fill=NAVY)
    base = shp(gs, MSO_SHAPE.ROUNDED_RECTANGLE, *P(0.30, 0.86), 0.40 * W, 0.08 * h, fill=NAVY)
    base.adjustments[0] = 0.5
    return g


def monitor_width(h):
    return h * 1.25


def _film(gs, x, y, W, h, fill=FILM):
    f = shp(gs, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, W, h, fill=fill, line=NAVY, lw=0.75)
    f.adjustments[0] = 0.08
    return f


def xray(shapes, x, y, h, name='X-ray thumbnail'):
    W = h * 0.82
    g = group(shapes, name)
    gs = g.shapes
    P = lambda fx, fy: (x + fx * W, y + fy * h)
    _film(gs, x, y, W, h)
    shp(gs, MSO_SHAPE.RECTANGLE, *P(0.47, 0.13), 0.06 * W, 0.76 * h, fill=FILM_LIGHT)
    rl = max(0.75, h * 1.6)
    line(gs, *P(0.45, 0.17), *P(0.18, 0.14), color=FILM_LIGHT, lw=rl)
    line(gs, *P(0.55, 0.17), *P(0.82, 0.14), color=FILM_LIGHT, lw=rl)
    for k in range(4):
        fy = 0.26 + k * 0.13
        poly(gs, [P(0.45, fy), P(0.30, fy + 0.03), P(0.18, fy + 0.10)], line=FILM_LIGHT, lw=rl)
        poly(gs, [P(0.55, fy), P(0.70, fy + 0.03), P(0.82, fy + 0.10)], line=FILM_LIGHT, lw=rl)
    return g


def ct_scan(shapes, x, y, h, name='CT / MRI thumbnail'):
    W = h * 0.82
    g = group(shapes, name)
    gs = g.shapes
    P = lambda fx, fy: (x + fx * W, y + fy * h)
    _film(gs, x, y, W, h)
    d = 0.62 * W
    cx, cy = P(0.5, 0.5)
    shp(gs, MSO_SHAPE.OVAL, cx - d / 2, cy - d * 0.6, d, d * 1.2, fill='7F97B2', line=FILM_LIGHT, lw=max(1.0, h * 3.0))
    line(gs, cx, cy - d * 0.45, cx, cy + d * 0.45, color=FILM, lw=max(0.5, h * 1.2))
    for sx in (-1, 1):
        dd = d * 0.22
        shp(gs, MSO_SHAPE.OVAL, cx + sx * d * 0.2 - dd / 2, cy - dd / 2, dd, dd * 1.3, fill=FILM)
    return g


def ultrasound(shapes, x, y, h, name='Ultrasound thumbnail'):
    W = h * 0.82
    g = group(shapes, name)
    gs = g.shapes
    P = lambda fx, fy: (x + fx * W, y + fy * h)
    _film(gs, x, y, W, h)
    ax, ay = P(0.5, 0.12)
    r = 0.74 * h
    pts = [(ax, ay)]
    for k in range(13):
        a = math.radians(58 + k * (64 / 12))
        pts.append((ax + r * math.cos(a), ay + r * math.sin(a)))
    poly(gs, pts, line=None, fill='9FB4CB', closed=True)
    bx, by = P(0.5, 0.60)
    shp(gs, MSO_SHAPE.OVAL, bx - 0.13 * W, by - 0.07 * h, 0.26 * W, 0.15 * h, fill=FILM)
    return g


def ecg(shapes, x, y, h, name='ECG strip thumbnail'):
    W = h * 1.3
    g = group(shapes, name)
    gs = g.shapes
    P = lambda fx, fy: (x + fx * W, y + fy * h)
    paper = shp(gs, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, W, h, fill=WHITE, line=NAVY, lw=0.75)
    paper.adjustments[0] = 0.08
    for fx in (0.25, 0.5, 0.75):
        line(gs, *P(fx, 0.06), *P(fx, 0.94), color=GRID, lw=0.5)
    for fy in (0.33, 0.66):
        line(gs, *P(0.04, fy), *P(0.96, fy), color=GRID, lw=0.5)
    trace = [(0.04, 0.58), (0.15, 0.58), (0.19, 0.50), (0.23, 0.58), (0.29, 0.58), (0.32, 0.66), (0.36, 0.18),
             (0.40, 0.82), (0.44, 0.58), (0.54, 0.58), (0.59, 0.45), (0.65, 0.58), (0.76, 0.58), (0.79, 0.66),
             (0.83, 0.20), (0.87, 0.80), (0.91, 0.58), (0.96, 0.58)]
    poly(gs, [P(a, b) for a, b in trace], line=BLUE, lw=max(0.75, h * 2.6))
    return g


def ecg_width(h):
    return h * 1.3


def lab_report(shapes, x, y, h, name='Laboratory report thumbnail'):
    W = h * 0.82
    g = group(shapes, name)
    gs = g.shapes
    P = lambda fx, fy: (x + fx * W, y + fy * h)
    shp(gs, MSO_SHAPE.FOLDED_CORNER, x, y, W, h, fill=WHITE, line=NAVY, lw=0.75)
    shp(gs, MSO_SHAPE.RECTANGLE, *P(0.12, 0.09), 0.56 * W, 0.08 * h, fill=BLUE)
    for fy, fw in ((0.25, 0.74), (0.34, 0.60), (0.43, 0.70)):
        shp(gs, MSO_SHAPE.RECTANGLE, *P(0.12, fy), fw * W, 0.045 * h, fill=SLATE)
    base = 0.78
    for fx, top, col in ((0.14, 0.60, NAVY), (0.27, 0.52, BLUE), (0.40, 0.64, NAVY)):
        shp(gs, MSO_SHAPE.RECTANGLE, *P(fx, top), 0.09 * W, (base - top) * h, fill=col)
    line(gs, *P(0.12, base), *P(0.56, base), color=MUTED, lw=0.5)
    return g


def doc_width(h):
    return h * 0.82


def clipboard(shapes, x, y, h, name='Intake card illustration'):
    W = h * 0.82
    g = group(shapes, name)
    gs = g.shapes
    P = lambda fx, fy: (x + fx * W, y + fy * h)
    lw = _lw(h)
    b = shp(gs, MSO_SHAPE.ROUNDED_RECTANGLE, *P(0.04, 0.07), 0.92 * W, 0.93 * h, fill=WHITE, line=NAVY, lw=lw)
    b.adjustments[0] = 0.08
    clip = shp(gs, MSO_SHAPE.ROUNDED_RECTANGLE, *P(0.30, 0.0), 0.40 * W, 0.14 * h, fill=NAVY)
    clip.adjustments[0] = 0.3
    hs = 0.17 * h
    hx, hy = P(0.24, 0.33)
    shp(gs, MSO_SHAPE.HEART, hx - hs / 2, hy - hs / 2, hs, hs * 0.9, fill=BLUE)
    poly(gs, [P(0.40, 0.335), P(0.50, 0.335), P(0.55, 0.23), P(0.61, 0.44), P(0.66, 0.335), P(0.86, 0.335)], line=BLUE,
         lw=max(0.75, lw * 0.9))
    for fy, fw in ((0.56, 0.68), (0.69, 0.56), (0.82, 0.62)):
        shp(gs, MSO_SHAPE.RECTANGLE, *P(0.16, fy), fw * W, 0.05 * h, fill=SLATE)
    return g


def encounter_card(shapes, x, y, h, name='Encounter record illustration'):
    W = h * 1.2
    g = group(shapes, name)
    gs = g.shapes
    P = lambda fx, fy: (x + fx * W, y + fy * h)
    lw = _lw(h)
    c = shp(gs, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, W, h, fill=WHITE, line=NAVY, lw=lw)
    c.adjustments[0] = 0.08
    hdr = shp(gs, MSO_SHAPE.RECTANGLE, *P(0.0, 0.0), W, 0.22 * h, fill=BLUE)
    no_effects(hdr)
    d = 0.13 * h
    px, py = P(0.12, 0.11)
    shp(gs, MSO_SHAPE.OVAL, px - d / 2, py - d / 2, d, d, fill=WHITE)
    shp(gs, MSO_SHAPE.RECTANGLE, *P(0.22, 0.085), 0.42 * W, 0.05 * h, fill=WHITE)
    hs = 0.14 * h
    hx, hy = P(0.13, 0.40)
    shp(gs, MSO_SHAPE.HEART, hx - hs / 2, hy - hs / 2, hs, hs * 0.9, fill=BLUE)
    poly(gs, [P(0.24, 0.40), P(0.38, 0.40), P(0.43, 0.31), P(0.49, 0.49), P(0.54, 0.40), P(0.86, 0.40)], line=BLUE,
         lw=max(0.75, lw * 0.8))
    for fy, fw in ((0.60, 0.62), (0.76, 0.48)):
        shp(gs, MSO_SHAPE.RECTANGLE, *P(0.08, fy), fw * W, 0.06 * h, fill=SLATE)
    cd = 0.15 * h
    cx, cy = P(0.86, 0.79)
    shp(gs, MSO_SHAPE.OVAL, cx - cd / 2, cy - cd / 2, cd, cd, fill=WHITE, line=MUTED, lw=0.75)
    line(gs, cx, cy, cx, cy - cd * 0.32, color=MUTED, lw=0.75)
    line(gs, cx, cy, cx + cd * 0.26, cy, color=MUTED, lw=0.75)
    return g


def encounter_width(h):
    return h * 1.2


def rx_card(shapes, x, y, h, state='draft', name=None):
    """Prescription sheet: state 'draft' (dashed, amber DRAFT tag) or 'authorised' (signature + green check)."""
    W = h * 0.82
    g = group(shapes, name or f'Prescription {state} illustration')
    gs = g.shapes
    P = lambda fx, fy: (x + fx * W, y + fy * h)
    lw = _lw(h)
    draft = state == 'draft'
    pg = shp(gs, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, W, h, fill=WHITE, line=AMBER if draft else NAVY, lw=lw,
             dash=DASH if draft else None)
    pg.adjustments[0] = 0.07
    text(gs, *P(0.08, 0.02), 0.62 * W, 0.40 * h, ['Rx'], size=max(7, round(h * 20)), bold=True, italic=True,
         color=NAVY, font='Times New Roman', align='l', anchor='t')
    for fy, fw in ((0.42, 0.72), (0.53, 0.60), (0.64, 0.66)):
        shp(gs, MSO_SHAPE.RECTANGLE, *P(0.12, fy), fw * W, 0.045 * h, fill=SLATE)
    if draft:
        t = pill(gs, *P(0.07, 0.74), 0.86 * W, 0.23 * h, 'DRAFT', AMBER_BG, AMBER, AMBER_TX, size=5, ins=0.0)
        t.name = 'Draft tag'
    else:
        poly(gs, [P(0.12, 0.86), P(0.20, 0.78), P(0.26, 0.88), P(0.33, 0.79), P(0.40, 0.87), P(0.52, 0.82)],
             line=NAVY, lw=max(0.75, lw * 0.8))
        line(gs, *P(0.10, 0.91), *P(0.62, 0.91), color=MUTED, lw=0.5)
        badge(gs, *P(0.80, 0.82), 0.34 * h, 'ok')
    return g


def clock(shapes, x, y, h, color=NAVY, name='Timer illustration'):
    g = group(shapes, name)
    gs = g.shapes
    lw = _lw(h, k=3.0)
    d = h * 0.86
    cx, cy = x + h * 0.5, y + h * 0.14 + d / 2
    shp(gs, MSO_SHAPE.ROUNDED_RECTANGLE, cx - h * 0.09, y, h * 0.18, h * 0.12, fill=color)
    shp(gs, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, fill=WHITE, line=color, lw=lw)
    line(gs, cx, cy, cx, cy - d * 0.32, color=color, lw=lw)
    line(gs, cx, cy, cx + d * 0.24, cy + d * 0.10, color=BLUE, lw=lw)
    dd = d * 0.10
    shp(gs, MSO_SHAPE.OVAL, cx - dd / 2, cy - dd / 2, dd, dd, fill=color)
    return g


def hub(shapes, cx, cy, w, label='JIVA', size=12, sub=None, name='JIVA hub'):
    h = w * 0.87
    s = shp(shapes, MSO_SHAPE.HEXAGON, cx - w / 2, cy - h / 2, w, h, fill=NAVY, line=WHITE, lw=1.25, shadow=True,
            name=name)
    paras = [[(label, {'size': size, 'bold': True, 'color': WHITE})]]
    if sub:
        paras.append([(sub, {'size': max(6, size - 5), 'bold': False, 'color': 'D6E2EF'})])
    fill_text(s.text_frame, paras, ins=(0.02, 0.0, 0.02, 0.0))
    return s


def database(shapes, x, y, w, h, fill=WHITE, edge=NAVY, name='Data store'):
    s = shp(shapes, MSO_SHAPE.CAN, x, y, w, h, fill=fill, line=edge, lw=1.0, name=name)
    s.adjustments[0] = 0.18
    return s


def shield(shapes, x, y, h, fill=NAVY, name='Shield illustration'):
    W = h * 0.84
    g = group(shapes, name)
    gs = g.shapes
    P = lambda fx, fy: (x + fx * W, y + fy * h)
    poly(gs, [P(0.5, 0.0), P(0.96, 0.14), P(0.93, 0.52), P(0.78, 0.78), P(0.5, 1.0), P(0.22, 0.78), P(0.07, 0.52),
              P(0.04, 0.14)], line=None, fill=fill, closed=True)
    poly(gs, [P(0.28, 0.50), P(0.44, 0.66), P(0.73, 0.34)], line=WHITE, lw=max(1.0, h * 4.0))
    return g


def doc_page(shapes, x, y, h, name='Clinical note illustration'):
    W = h * 0.82
    g = group(shapes, name)
    gs = g.shapes
    P = lambda fx, fy: (x + fx * W, y + fy * h)
    lw = _lw(h)
    pg = shp(gs, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, W, h, fill=WHITE, line=NAVY, lw=lw)
    pg.adjustments[0] = 0.07
    shp(gs, MSO_SHAPE.RECTANGLE, *P(0.14, 0.12), 0.50 * W, 0.08 * h, fill=NAVY)
    for fy, fw in ((0.30, 0.72), (0.42, 0.64), (0.54, 0.70), (0.66, 0.52)):
        shp(gs, MSO_SHAPE.RECTANGLE, *P(0.14, fy), fw * W, 0.045 * h, fill=SLATE)
    poly(gs, [P(0.52, 0.92), P(0.86, 0.58), P(0.94, 0.66), P(0.60, 1.0), P(0.50, 1.02)], line=None, fill=BLUE,
         closed=True)
    return g


def cloud(shapes, x, y, w, h, fill=WHITE, edge=BLUE, lw=1.0, name='Cloud illustration'):
    return shp(shapes, MSO_SHAPE.CLOUD, x, y, w, h, fill=fill, line=edge, lw=lw, name=name)


def coin(shapes, x, y, h, name='Cost illustration'):
    g = group(shapes, name)
    gs = g.shapes
    lw = _lw(h)
    shp(gs, MSO_SHAPE.OVAL, x + 0.10 * h, y + 0.22 * h, 0.78 * h, 0.78 * h, fill=TINT, line=NAVY, lw=lw)
    c = shp(gs, MSO_SHAPE.OVAL, x, y, 0.78 * h, 0.78 * h, fill=WHITE, line=NAVY, lw=lw)
    fill_text(c.text_frame, ['₹'], size=max(7, round(h * 26)), bold=True, color=NAVY, ins=(0, 0, 0, 0))
    return g
