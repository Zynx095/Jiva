"""Build JIVA_IDE_BOOTCAMP_FORMAT_VISUAL_ENHANCED.pptx from JIVA_IDE_BOOTCAMP_FORMAT_FINAL.pptx.

The input deck is opened read-only (python-pptx loads it into memory) and never written. Every original slide keeps
its order, layout, title, footer bar and slide number; four continuation slides in the same layout are inserted
inside the sections they belong to. All diagrams are native, editable PowerPoint shapes.

usage: python -I build_deck.py <input.pptx> <output.pptx>
"""
import copy
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.stdout.reconfigure(encoding='utf-8')

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Pt  # noqa: E402

import pptlib as L  # noqa: E402
from pptlib import (AMBER, AMBER_BG, AMBER_TINT, AMBER_TX, BLUE, DASH, E, GREEN, GREEN_BG, GREEN_TX, GREY,  # noqa: E402
                    GREY_BG, GREY_TX, MUTED, NAVY, RED, RED_BG, RED_TX, SLATE, SLATE_BG, TEXT, WHITE)

SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)
S = list(prs.slides)  # the six original slides, in order


# ---------------------------------------------------------------- generic helpers
def find(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(f'{name!r} not found')


def remove(slide, name):
    el = find(slide, name)._element
    el.getparent().remove(el)


def replace_in_runs(shape, old, new):
    """Replace text inside a single run, keeping that run's formatting. Fails loudly if not found exactly once."""
    hits = [r for p in shape.text_frame.paragraphs for r in p.runs if old in r.text]
    if len(hits) != 1:
        raise ValueError(f'{old!r}: {len(hits)} runs in {shape.name!r}')
    hits[0].text = hits[0].text.replace(old, new)


def clone_frame_slide(src_slide, title):
    """New slide in the template's 'Title and Content' layout with the source slide's title style, footer bar and
    slide-number field (copied element for element), so it is indistinguishable in format from the originals."""
    new = prs.slides.add_slide(src_slide.slide_layout)
    for ph in list(new.placeholders):
        ph._element.getparent().remove(ph._element)
    tree = new.shapes._spTree
    for name in ('Rectangle 9', 'Title 1', 'Slide Number Placeholder 5'):
        tree.append(copy.deepcopy(find(src_slide, name)._element))
    runs = find(new, 'Title 1').text_frame.paragraphs[0].runs
    runs[0].text = title
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    return new


def move_slide(slide, index):
    lst = prs.slides._sldIdLst
    for sldId in list(lst):
        if prs.part.related_part(sldId.rId) is slide.part:
            lst.remove(sldId)
            lst.insert(index, sldId)
            return
    raise KeyError('slide not found')


def panel(sh, x, y, w, h, line=BLUE, lw=1.5, radius=0.42, name='Panel'):
    return L.card(sh, x, y, w, h, fill=WHITE, line=line, lw=lw, radius=radius, name=name)


def heading(sh, x, y, w, label, name='Panel heading', size=13):
    return L.text(sh, x, y, w, 0.30, [label], size=size, bold=True, color=NAVY, align='l', anchor='t', name=name)


def step_card(sh, x, y, w, h, n, title, desc, top=0.64, fill=WHITE, line=SLATE, lw=1.0, name=None, desc_size=8.5):
    c = L.card(sh, x, y, w, h, fill=fill, line=line, lw=lw, radius=0.14, name=name or f'Step {n}: {title}')
    paras = [[(title, {'size': 11, 'bold': True, 'color': NAVY})]]
    if desc:
        paras.append(([(desc, {'size': desc_size, 'color': MUTED})], {'space_before': 2}))
    L.fill_text(c.text_frame, paras, anchor='t', ins=(0.08, top, 0.08, 0.04))
    L.num_badge(sh, x + 0.06, y + 0.02, 0.32, n)
    return c


def hcenter(x, w, total):
    return x + (w - total) / 2


# ================================================================ SLIDE 1 — title page: care-loop summary
def slide1_title():
    s = S[0]
    sh = s.shapes
    g = L.group(sh, 'Shared care state illustration')
    gs = g.shapes
    cx, cy, R, d = 9.20, 4.15, 1.50, 0.80
    import math
    nodes = [
        (-90, 'Ambulance crew', lambda x, y: L.ambulance(gs, x - L.ambulance_width(0.30) / 2, y - 0.15, 0.30)),
        (-18, 'Hospital', lambda x, y: L.hospital(gs, x - L.hospital_width(0.42) / 2, y - 0.22, 0.42)),
        (54, 'Doctor', lambda x, y: L.doctor(gs, x - 0.22, y - 0.24, 0.44)),
        (126, 'Patient', lambda x, y: L.patient(gs, x - 0.22, y - 0.24, 0.44)),
        (198, 'Management', lambda x, y: L.monitor(gs, x - L.monitor_width(0.36) / 2, y - 0.18, 0.36)),
    ]
    pts = []
    for ang, label, draw in nodes:
        a = math.radians(ang)
        pts.append((cx + R * math.cos(a), cy + R * math.sin(a), label, draw))
    for x, y, _, _ in pts:  # spokes first so circles sit on top
        dx, dy = x - cx, y - cy
        n = math.hypot(dx, dy)
        stop = d / 2 + (0.28 if dy / n < -0.9 else 0.02)  # the top node's label sits under its circle
        L.line(gs, cx + dx / n * 0.52, cy + dy / n * 0.47, x - dx / n * stop, y - dy / n * stop, color=BLUE, lw=1.25)
    for x, y, label, draw in pts:
        L.shp(gs, MSO_SHAPE.OVAL, x - d / 2, y - d / 2, d, d, fill=WHITE, line=SLATE, lw=1.0, shadow=True,
              name=f'{label} node')
        draw(x, y)
        L.text(gs, x - 0.75, y + d / 2 + 0.02, 1.5, 0.22, [label], size=9, bold=True, color=NAVY)
    L.hub(gs, cx, cy, 1.10, label='JIVA', size=14, sub='shared care state')


# ================================================================ SLIDE 2 — proposed solution: icons, pills, problem/solution visuals
def slide2_solution():
    s = S[1]
    sh = s.shapes
    cards = ['Rounded Rectangle 15364', 'Rounded Rectangle 15365', 'Rounded Rectangle 15366', 'Rounded Rectangle 15367',
             'Rounded Rectangle 15368']
    for nm in cards:
        tf = find(s, nm).text_frame
        bp = tf._txBody.bodyPr
        bp.set('anchor', 't')
        bp.set('tIns', str(E(0.50)))
    X = [0.80, 3.15, 5.50, 7.85, 10.20]
    top, W = 1.75, 2.05
    # 1 Emergency intake: ambulance + intake card
    tw = L.ambulance_width(0.34) + 0.08 + L.doc_width(0.42)
    x0 = hcenter(X[0], W, tw)
    L.ambulance(sh, x0, top + 0.13, 0.34)
    L.clipboard(sh, x0 + L.ambulance_width(0.34) + 0.08, top + 0.08, 0.42)
    # 2 Feasibility gate: eligible vs not eligible
    hw = L.hospital_width(0.38)
    x0 = hcenter(X[1], W, hw * 2 + 0.22)
    L.hospital(sh, x0, top + 0.10, 0.38, status='ok', name='Eligible hospital')
    L.hospital(sh, x0 + hw + 0.22, top + 0.10, 0.38, status='no', muted=True, name='Not eligible hospital')
    # 3 Acceptance SLA: timer
    L.clock(sh, hcenter(X[2], W, 0.42), top + 0.07, 0.42)
    # 4 Explicit response: hospital + the four answers as pills (replaces the bracketed text)
    L.hospital(sh, hcenter(X[3], W, L.hospital_width(0.38)), top + 0.10, 0.38, name='Answering hospital')
    replace_in_runs(find(s, cards[3]), 'Mandatory active confirmation:\n[ACCEPTED]  [LIMITED]\n[REJECTED]  [UNAVAILABLE]\nSilence expires to UNKNOWN',
                    'Mandatory active confirmation:')
    pw, ph = 0.86, 0.19
    px = hcenter(X[3], W, pw * 2 + 0.07)
    for i, (label, kind) in enumerate((('ACCEPTED', 'done'), ('LIMITED', 'done'), ('REJECTED', 'no'), ('UNAVAILABLE', 'busy'))):
        f, b, c = {'done': (GREEN_BG, GREEN, GREEN_TX), 'no': (RED_BG, RED, RED_TX), 'busy': (AMBER_BG, AMBER, AMBER_TX)}[kind]
        L.pill(sh, px + (i % 2) * (pw + 0.07), top + 1.16 + (i // 2) * (ph + 0.04), pw, ph, label, f, b, c, size=7.5,
               name=f'Answer {label}')
    L.text(sh, X[3] + 0.05, top + 1.60, W - 0.10, 0.18, ['Silence expires to UNKNOWN'], size=9.5, color=MUTED,
           name='Expiry note', anchor='m')
    # 5 Dynamic reroute: lost destination -> accepting alternative
    hw = L.hospital_width(0.36)
    x0 = hcenter(X[4], W, hw * 2 + 0.34)
    L.hospital(sh, x0, top + 0.11, 0.36, status='no', name='Lost destination')
    L.hospital(sh, x0 + hw + 0.34, top + 0.11, 0.36, status='ok', name='Alternative destination')
    L.arrow(sh, x0 + hw + 0.04, top + 0.33, x0 + hw + 0.30, top + 0.33, color=BLUE, lw=1.5)

    # Problem vs solution: same three parties, fragmented vs connected
    def parties(gs, x, y, connected):
        amb = (x + 0.02, y + 0.06, 0.30)
        hos = (x + 0.86, y + 0.0, 0.40)
        doc = (x + 0.44, y + 0.92, 0.42)
        A = (amb[0] + L.ambulance_width(0.30) / 2, amb[1] + 0.15)
        H = (hos[0] + L.hospital_width(0.40) / 2, hos[1] + 0.22)
        D = (doc[0] + 0.21, doc[1] + 0.22)
        if connected:
            hub = (x + 0.66, y + 0.56)
            for p in (A, H, D):
                L.line(gs, hub[0], hub[1], p[0], p[1], color=BLUE, lw=1.5)
        else:
            for p, q in ((A, H), (A, D), (H, D)):
                L.line(gs, p[0], p[1], q[0], q[1], color=RED, lw=1.25, dash=DASH)
                L.badge(gs, (p[0] + q[0]) / 2, (p[1] + q[1]) / 2, 0.17, 'no')
        L.ambulance(gs, *amb)
        L.hospital(gs, *hos)
        L.doctor(gs, *doc)
        if connected:
            L.hub(gs, hub[0], hub[1], 0.40, label='JIVA', size=7)

    for nm, x, connected, caption, col in (('Rounded Rectangle 15373', 0.60, False, 'Isolated calls, no shared picture', RED_TX),
                                           ('Rounded Rectangle 15374', 6.78, True, 'One shared, timestamped state', GREEN_TX)):
        box = find(s, nm)
        bp = box.text_frame._txBody.bodyPr
        bp.set('anchor', 't')
        bp.set('tIns', str(E(0.16)))
        for p in box.text_frame.paragraphs[1:]:
            p._p.get_or_add_pPr().set('marL', str(E(1.50)))
        g = L.group(sh, 'Fragmented care illustration' if not connected else 'Connected care illustration')
        parties(g.shapes, x + 0.20, 4.72, connected)
        L.text(g.shapes, x + 0.10, 6.22, 1.50, 0.34, [caption], size=8, italic=True, bold=True, color=col)

    # accuracy: ports and app count; SHA-256 wording (decision records are hashed; there is no tamper-evident ledger)
    green = find(s, 'Rounded Rectangle 15374')
    replace_in_runs(green, 'Lightweight web overlay across 4 role-scoped apps (:5174-:5177)',
                    'Lightweight web overlay across 5 role-scoped apps (:5173–:5177)')
    replace_in_runs(green, 'SHA-256 event audits', 'SHA-256-hashed decision audits')


# ================================================================ NEW — the emergency-to-care journey (flagship)
def journey_slide():
    s = clone_frame_slide(S[2], 'THE EMERGENCY-TO-CARE JOURNEY')
    sh = s.shapes
    panel(sh, 0.60, 1.20, 12.13, 5.62, name='Journey panel')
    heading(sh, 0.85, 1.27, 11.6, 'FLAGSHIP WORKFLOW: from the emergency call to clinician-authorised care')
    # lanes
    L.card(sh, 0.80, 1.66, 11.73, 2.30, fill='F4F8FC', line=None, radius=0.14, shadow=False, name='Lane: emergency coordination')
    L.card(sh, 0.80, 4.14, 11.73, 2.30, fill=AMBER_TINT, line=None, radius=0.14, shadow=False, name='Lane: clinical care')
    L.text(sh, 0.98, 1.71, 2.18, 0.18, ['EMERGENCY COORDINATION'], size=9, bold=True, color=NAVY, align='l', anchor='m')
    L.status_pill(sh, 3.22, 1.71, 2.35, 0.18, '✓ IMPLEMENTED · VERIFIED LOCALLY', 'done', size=7.5)
    L.text(sh, 0.98, 4.19, 1.32, 0.18, ['CLINICAL CARE'], size=9, bold=True, color=NAVY, align='l', anchor='m')
    L.status_pill(sh, 2.36, 4.19, 2.75, 0.18, 'WORKING PROTOTYPE · IN DEVELOPMENT', 'dev', size=7.5)

    Wc = [1.80, 1.80, 2.55, 1.95, 1.95]
    gap = (11.63 - sum(Wc)) / 4
    X = [0.85]
    for w in Wc[:-1]:
        X.append(X[-1] + w + gap)
    y1, y2, H = 2.04, 4.54, 1.60
    mid1, mid2 = y1 + H / 2, y2 + H / 2

    # 1 Emergency intake
    step_card(sh, X[0], y1, Wc[0], H, 1, 'Emergency intake', 'Crew reports symptoms, severity and location; vital signs typed in by hand.')
    tw = L.ambulance_width(0.34) + 0.07 + L.doc_width(0.44)
    x0 = hcenter(X[0], Wc[0], tw)
    L.ambulance(sh, x0, y1 + 0.17, 0.34)
    L.clipboard(sh, x0 + L.ambulance_width(0.34) + 0.07, y1 + 0.10, 0.44)
    # 2 Hospital referral: fan-out to eligible hospitals only
    step_card(sh, X[1], y1, Wc[1], H, 2, 'Hospital referral', 'Eligible hospitals (capability, status) get time-limited referrals in parallel.')
    hw = L.hospital_width(0.30)
    x0 = hcenter(X[1], Wc[1], hw * 3 + 0.16)
    bus_y = y1 + 0.16
    L.line(sh, x0 + hw / 2, bus_y, x0 + hw * 1.5 + 0.08, bus_y, color=BLUE, lw=1.25)
    for i in range(3):
        hx = x0 + i * (hw + 0.08)
        muted = i == 2
        if not muted:
            L.arrow(sh, hx + hw / 2, bus_y, hx + hw / 2, y1 + 0.27, color=BLUE, lw=1.25, size='sm')
        L.hospital(sh, hx, y1 + 0.28, 0.30, status='no' if muted else None, muted=muted,
                   name='Not eligible hospital' if muted else 'Referred hospital')
    # 3 Hospital response: explicit answers branch from the hospital
    step_card(sh, X[2], y1, Wc[2], H, 3, 'Hospital response', 'Answers are explicit and time-limited; silence is never a yes.', top=0.98)
    hx, hy = X[2] + 0.16, y1 + 0.22
    L.hospital(sh, hx, hy, 0.46, name='Responding hospital')
    chips = (('Accepts (or limited)', 'ok'), ('Rejects', 'no'), ('Busy: unavailable', 'busy'), ('No answer: expires', 'expired'))
    cx0 = X[2] + 1.02
    for i, (label, kind) in enumerate(chips):
        cy = y1 + 0.10 + i * 0.205
        f, b, c = {'ok': (GREEN_BG, GREEN, GREEN_TX), 'no': (RED_BG, RED, RED_TX), 'busy': (AMBER_BG, AMBER, AMBER_TX),
                   'expired': (GREY_BG, GREY, GREY_TX)}[kind]
        L.line(sh, hx + L.hospital_width(0.46) + 0.02, hy + 0.23, cx0, cy + 0.09, color=b, lw=1.0)
        L.pill(sh, cx0, cy, 1.42, 0.18, label, f, b, c, size=7.5, name=f'Response: {label}')
    # 4 Dynamic rerouting
    step_card(sh, X[3], y1, Wc[3], H, 4, 'Dynamic rerouting', 'Destination rejects, turns unavailable or lapses: JIVA re-checks hospitals that accepted.')
    hw = L.hospital_width(0.38)
    x0 = hcenter(X[3], Wc[3], hw * 2 + 0.36)
    L.hospital(sh, x0, y1 + 0.14, 0.38, status='no', name='Lost destination')
    L.hospital(sh, x0 + hw + 0.36, y1 + 0.14, 0.38, status='ok', name='Accepting alternative')
    L.arrow(sh, x0 + hw + 0.05, y1 + 0.38, x0 + hw + 0.31, y1 + 0.38, color=BLUE, lw=1.5)
    # 5 Referral accepted (highlighted)
    step_card(sh, X[4], y1, Wc[4], H, 5, 'Referral accepted', 'Best eligible hospital that explicitly accepted; ETA only orders the choice.',
              fill='F0FDF4', line=GREEN, lw=1.75)
    aw = L.ambulance_width(0.30)
    hw = L.hospital_width(0.42)
    x0 = hcenter(X[4], Wc[4], aw + 0.30 + hw)
    L.ambulance(sh, x0, y1 + 0.26, 0.30)
    L.arrow(sh, x0 + aw + 0.04, y1 + 0.40, x0 + aw + 0.26, y1 + 0.40, color=GREEN, lw=1.5)
    L.hospital(sh, x0 + aw + 0.30, y1 + 0.12, 0.42, status='ok', name='Destination hospital')

    # row-1 flow
    L.arrow(sh, X[0] + Wc[0] + 0.04, mid1, X[1] - 0.04, mid1, color=BLUE, lw=1.75)
    L.arrow(sh, X[1] + Wc[1] + 0.04, mid1, X[2] - 0.04, mid1, color=BLUE, lw=1.75)
    L.arrow(sh, X[2] + Wc[2] + 0.04, mid1, X[3] - 0.04, mid1, color=RED, lw=1.75, name='Rejected, busy or expired')
    L.arrow(sh, X[3] + Wc[3] + 0.04, mid1, X[4] - 0.04, mid1, color=BLUE, lw=1.75, name='Another hospital accepted')
    ax = X[2] + Wc[2] - 0.30
    bx = X[4] + Wc[4] / 2
    L.poly(sh, [(ax, y1 - 0.01), (ax, 1.83), (bx, 1.83), (bx, y1 - 0.01)], line=GREEN, lw=1.75, tail='triangle',
           name='Accepted path')
    lab = L.text(sh, X[3] + 0.05, 1.74, Wc[3] - 0.10, 0.155, ['accepted / limited → destination'], size=8, bold=True,
                 color=GREEN_TX, name='Accepted path label', anchor='m')
    lab.fill.solid()
    lab.fill.fore_color.rgb = L.rgb('F4F8FC')
    lx1, lx2 = X[3] + 0.45, X[2] + Wc[2] - 0.55
    L.poly(sh, [(lx1, y1 + H + 0.01), (lx1, 3.78), (lx2, 3.78), (lx2, y1 + H + 0.01)], line=GREY, lw=1.25, dash=DASH,
           tail='triangle', name='Awaiting responses loop')
    L.text(sh, lx2 - 0.55, 3.79, (lx1 - lx2) + 1.10, 0.16, ['no hospital has accepted yet → awaiting responses'],
           size=7.5, italic=True, color=GREY_TX, name='Awaiting responses label')

    # U-turn: ambulance arrives at the accepted hospital
    ux = X[4] + Wc[4] / 2
    L.arrow(sh, ux, y1 + H + 0.02, ux, y2 - 0.02, color=BLUE, lw=1.75, name='Ambulance arrives')
    L.text(sh, ux - 1.40, 3.98, 1.32, 0.16, ['ambulance arrives'], size=8, italic=True, bold=True, color=NAVY,
           align='r', name='Arrival label')

    # row 2 (right to left): 6 under 5, 7 under 4, 8 under 3, 9 under 2, 10 under 1
    step_card(sh, X[4], y2, Wc[4], H, 6, 'Arrival & handover', 'Hospital records arrival and handover; the encounter opens for its care team.')
    aw = L.ambulance_width(0.30)
    hw = L.hospital_width(0.44)
    x0 = hcenter(X[4], Wc[4], aw + 0.06 + hw)
    L.hospital(sh, x0 + aw + 0.06, y2 + 0.12, 0.44, name='Receiving hospital')
    L.ambulance(sh, x0, y2 + 0.26, 0.30)

    step_card(sh, X[3], y2, Wc[3], H, 7, 'Doctor workspace', "Assigned doctor reviews symptoms, timestamped vitals and the crew's handover.")
    ew = L.encounter_width(0.38)
    x0 = hcenter(X[3], Wc[3], 0.46 + 0.10 + ew)
    L.doctor(sh, x0, y2 + 0.10, 0.46)
    L.encounter_card(sh, x0 + 0.56, y2 + 0.16, 0.38)

    step_card(sh, X[2], y2, Wc[2], H, 8, 'Clinical documents', 'Typed and versioned; never auto-interpreted. Demo samples are synthetic.', top=0.84)
    th = 0.44
    items = [(L.xray, L.doc_width(th), 'X-ray'), (L.ecg, L.ecg_width(th), 'ECG'), (L.ultrasound, L.doc_width(th), 'Ultrasound'),
             (L.lab_report, L.doc_width(th), 'Lab report')]
    tot = sum(w for _, w, _ in items) + 0.16 * 3
    x0 = hcenter(X[2], Wc[2], tot)
    for fn, w, label in items:
        fn(sh, x0, y2 + 0.10, th)
        L.text(sh, x0 - 0.075, y2 + 0.10 + th + 0.03, w + 0.15, 0.15, [label], size=7, color=MUTED, name=f'{label} caption')
        x0 += w + 0.16

    step_card(sh, X[1], y2, Wc[1], H, 9, 'Clinical decision', 'The doctor assesses; a prescription stays a draft until explicitly authorised.')
    rw = L.doc_width(0.40)
    x0 = hcenter(X[1], Wc[1], 0.42 + 0.06 + rw + 0.22 + rw)
    L.doctor(sh, x0, y2 + 0.12, 0.42)
    L.rx_card(sh, x0 + 0.48, y2 + 0.13, 0.40, 'draft')
    L.arrow(sh, x0 + 0.48 + rw + 0.03, y2 + 0.33, x0 + 0.48 + rw + 0.19, y2 + 0.33, color=BLUE, lw=1.25, size='sm')
    L.rx_card(sh, x0 + 0.48 + rw + 0.22, y2 + 0.13, 0.40, 'authorised')

    step_card(sh, X[0], y2, Wc[0], H, 10, 'Shared visibility', 'Patient: authorised items · Doctor: full encounter · Management: status only.')
    mw = L.monitor_width(0.34)
    x0 = hcenter(X[0], Wc[0], 0.40 * 2 + mw + 0.16)
    L.patient(sh, x0, y2 + 0.12, 0.40)
    L.doctor(sh, x0 + 0.48, y2 + 0.12, 0.40)
    L.monitor(sh, x0 + 0.96, y2 + 0.15, 0.34)

    for a, b in ((4, 3), (3, 2), (2, 1), (1, 0)):
        L.arrow(sh, X[a] - 0.04, mid2, X[b] + Wc[b] + 0.04, mid2, color=BLUE, lw=1.75)

    L.text(sh, 0.85, 6.50, 11.63, 0.26,
           ['Vital signs are entered manually by the crew or clinician — no medical-device integration. JIVA never diagnoses '
            'or prescribes: clinicians decide and authorise. All demo data is synthetic.'],
           size=8, italic=True, color=MUTED, name='Journey footnote')
    return s


# ================================================================ NEW — hospital responses & rerouting
def responses_slide():
    s = clone_frame_slide(S[2], 'HOSPITAL RESPONSES & REROUTING')
    sh = s.shapes
    panel(sh, 0.60, 1.20, 12.13, 5.62, name='Responses panel')
    heading(sh, 0.85, 1.27, 11.6, 'CLOSED-LOOP ACCEPTANCE: a destination needs an explicit, current acceptance — nothing is assumed')

    # A — new emergency
    a = L.card(sh, 0.85, 2.45, 1.55, 1.50, radius=0.14, name='New emergency')
    L.fill_text(a.text_frame, [[('New emergency', {'size': 10, 'bold': True, 'color': NAVY})],
                               ([('crew intake; care needs assessed', {'size': 8, 'color': MUTED})], {'space_before': 2})],
                anchor='t', ins=(0.06, 0.66, 0.06, 0.04))
    L.ambulance(sh, hcenter(0.85, 1.55, L.ambulance_width(0.34)), 2.60, 0.34)
    # B — referrals
    b = L.card(sh, 2.75, 2.45, 1.75, 1.50, radius=0.14, name='Referrals sent')
    L.fill_text(b.text_frame, [[('Referrals sent', {'size': 10, 'bold': True, 'color': NAVY})],
                               ([('to every eligible hospital at once; each time-limited', {'size': 8, 'color': MUTED})], {'space_before': 2})],
                anchor='t', ins=(0.06, 0.66, 0.06, 0.04))
    hw = L.hospital_width(0.28)
    x0 = hcenter(2.75, 1.75, hw * 3 + 0.16)
    L.line(sh, x0 + hw / 2, 2.57, x0 + hw * 2.5 + 0.16, 2.57, color=BLUE, lw=1.25)
    for i in range(3):
        hx = x0 + i * (hw + 0.08)
        L.arrow(sh, hx + hw / 2, 2.57, hx + hw / 2, 2.67, color=BLUE, lw=1.25, size='sm')
        L.hospital(sh, hx, 2.68, 0.28, name='Referred hospital')
    L.arrow(sh, 2.42, 3.20, 2.73, 3.20, color=BLUE, lw=1.75)
    # C — decision
    dmd = L.shp(sh, MSO_SHAPE.FLOWCHART_DECISION, 4.85, 2.45, 1.60, 1.50, fill=WHITE, line=NAVY, lw=1.5, shadow=True,
                name='Hospital response decision')
    L.fill_text(dmd.text_frame, [[('Hospital', {'size': 10, 'bold': True, 'color': NAVY})],
                                 [('response', {'size': 10, 'bold': True, 'color': NAVY})]], ins=(0, 0, 0, 0))
    L.arrow(sh, 4.52, 3.20, 4.83, 3.20, color=BLUE, lw=1.75)
    # D — the four outcomes
    outs = [
        ('ok', 'ACCEPTED / LIMITED', 'explicit yes, valid until a stated time', GREEN, GREEN_BG, GREEN_TX),
        ('no', 'REJECTED', 'excluded for this case', RED, RED_BG, RED_TX),
        ('busy', 'UNAVAILABLE (busy)', 'not usable while unavailable', AMBER, AMBER_BG, AMBER_TX),
        ('expired', 'EXPIRED', 'no answer in time, or the acceptance lapsed', GREY, GREY_BG, GREY_TX),
    ]
    DX, DW, DH = 6.90, 2.45, 0.62
    DY = [1.78, 2.52, 3.26, 4.00]
    for (kind, title, desc, col, bg, tx), dy in zip(outs, DY):
        c = L.card(sh, DX, dy, DW, DH, fill=bg, line=col, lw=1.25, radius=0.10, name=f'Outcome {title}')
        L.fill_text(c.text_frame, [[(title, {'size': 9, 'bold': True, 'color': tx})],
                                   [(desc, {'size': 7.5, 'color': TEXT})]], anchor='m', align='l', ins=(0.62, 0.02, 0.06, 0.02))
        L.hospital(sh, DX + 0.10, dy + 0.10, 0.40, status=kind, muted=kind in ('no', 'expired'), name=f'{title} hospital')
        L.arrow(sh, 6.45, 3.20, DX - 0.03, dy + DH / 2, color=col, lw=1.5)
    # E — results
    EX, EW = 9.80, 2.70
    e1 = L.card(sh, EX, 1.78, EW, 0.95, fill='F0FDF4', line=GREEN, lw=1.75, radius=0.12, name='Destination confirmed')
    L.fill_text(e1.text_frame, [[('DESTINATION CONFIRMED', {'size': 9.5, 'bold': True, 'color': GREEN_TX})],
                                [('Best eligible hospital with a current acceptance; ETA only orders the choice. Route sent to the ambulance.', {'size': 7.5, 'color': TEXT})]],
                anchor='m', align='l', ins=(0.10, 0.04, 0.08, 0.04))
    e2 = L.card(sh, EX, 3.02, EW, 0.92, fill='F4F8FC', line=BLUE, lw=1.5, radius=0.12, name='Re-evaluate')
    L.fill_text(e2.text_frame, [[('RE-EVALUATE', {'size': 9.5, 'bold': True, 'color': BLUE})],
                                [('only when the current destination is lost:', {'size': 7.5, 'bold': True, 'color': NAVY})],
                                [('next eligible hospital with a current acceptance (existing routing rules)', {'size': 7.5, 'color': TEXT})]],
                anchor='m', align='l', ins=(0.10, 0.03, 0.08, 0.03))
    e3 = L.card(sh, EX, 4.20, EW, 0.68, fill=GREY_BG, line=GREY, lw=1.25, radius=0.12, name='No destination yet')
    L.fill_text(e3.text_frame, [[('NO DESTINATION YET', {'size': 9.5, 'bold': True, 'color': GREY_TX})],
                                [('Nobody has accepted: route withdrawn, awaiting responses.', {'size': 7.5, 'color': TEXT})]],
                anchor='m', align='l', ins=(0.10, 0.04, 0.08, 0.04))
    L.arrow(sh, DX + DW + 0.03, DY[0] + DH / 2, EX - 0.03, DY[0] + DH / 2, color=GREEN, lw=1.75)
    for (kind, *_r), dy in zip(outs[1:], DY[1:]):
        col = {'no': RED, 'busy': AMBER, 'expired': GREY}[kind]
        L.arrow(sh, DX + DW + 0.03, dy + DH / 2, EX - 0.03, 3.45, color=col, lw=1.5)
    L.arrow(sh, EX + 0.55, 3.00, EX + 0.55, 2.75, color=BLUE, lw=1.75, name='Found')
    L.text(sh, EX + 0.66, 2.75, 2.0, 0.24, ['found → destination changes, route recalculated'], size=7, italic=True,
           color=BLUE, align='l', anchor='m')
    L.arrow(sh, EX + 0.55, 3.96, EX + 0.55, 4.18, color=GREY, lw=1.5, name='None')
    L.text(sh, EX + 0.66, 3.95, 1.6, 0.24, ['none'], size=7, italic=True, color=GREY_TX, align='l', anchor='m')
    L.poly(sh, [(EX + EW / 2 + 0.6, 4.89), (EX + EW / 2 + 0.6, 5.02), (5.65, 5.02), (5.65, 3.97)], line=GREY, lw=1.25,
           dash=DASH, tail='triangle', name='Wait for outstanding referrals')
    L.text(sh, 6.0, 4.75, 3.6, 0.22, ['outstanding referrals can still be answered'], size=7.5, italic=True,
           color=GREY_TX, name='Loop label')

    # what the screens show (UI strings from the ambulance and management apps; synthetic example)
    L.line(sh, 0.85, 5.18, 12.48, 5.18, color=SLATE, lw=0.75)
    L.text(sh, 0.85, 5.22, 1.75, 0.20, ['WHAT THE SCREENS SHOW'], size=8.5, bold=True, color=NAVY, align='l', anchor='m')
    L.text(sh, 0.85, 5.43, 1.75, 0.15, ['example sequence, synthetic data'], size=7, italic=True, color=MUTED, align='l',
           anchor='m')
    L.ambulance(sh, 0.88, 5.66, 0.24)
    L.text(sh, 1.36, 5.60, 1.20, 0.38, ['Ambulance app'], size=8.5, bold=True, color=NAVY, align='l', anchor='m')
    L.monitor(sh, 0.90, 6.16, 0.28)
    L.text(sh, 1.36, 6.11, 1.20, 0.42, ['Management'], size=8.5, bold=True, color=NAVY, align='l', anchor='m')
    cols = [
        ('1 · Referrals out', 'Awaiting a hospital’s acceptance', 'Referral evidence: each hospital outstanding until it answers or expires', 'neutral'),
        ('2 · Hospital A accepts', 'Destination: Hospital A — route ready', 'A answered ACCEPTED, valid until a set time; case in transit', 'done'),
        ('3 · Hospital A reports unavailable', 'Rerouted to Hospital B (it had accepted) — reason shown', 'Destination changed, with the reason; A unavailable', 'busy'),
        ('Alternative · nobody has accepted', 'Destination withdrawn — waiting for another hospital to accept', 'Referral evidence shows who is still outstanding', 'plan'),
    ]
    CX0, CW, CG = 2.62, 2.38, 0.08
    for i, (hdr, amb, mgmt, kind) in enumerate(cols):
        x = CX0 + i * (CW + CG)
        f, bcol, tcol = {'neutral': (WHITE, SLATE, NAVY), 'done': (GREEN_BG, GREEN, GREEN_TX), 'busy': (AMBER_BG, AMBER, AMBER_TX),
                         'plan': (GREY_BG, GREY, GREY_TX)}[kind]
        L.text(sh, x, 5.24, CW, 0.22, [hdr], size=8, bold=True, color=tcol if kind != 'neutral' else BLUE, anchor='m')
        for y, body in ((5.50, amb), (6.04, mgmt)):
            c = L.card(sh, x, y, CW, 0.50, fill=WHITE if kind == 'neutral' else f, line=bcol, lw=1.0, radius=0.08,
                       shadow=False, dash=DASH if kind == 'plan' else None, name=f'{hdr}: {body[:20]}')
            L.fill_text(c.text_frame, [body], size=7.5, color=TEXT, ins=(0.06, 0.02, 0.06, 0.02))
        if i < 2:
            L.arrow(sh, x + CW + 0.005, 5.75, x + CW + CG - 0.005, 5.75, color=BLUE, lw=1.0, size='sm')
    return s


# ================================================================ SLIDE 3 — technical approach: system architecture diagram
def slide3_architecture():
    s = S[2]
    sh = s.shapes
    for nm in ('Rounded Rectangle 17412', 'Right Arrow 17413', 'Rounded Rectangle 17414', 'Right Arrow 17415',
               'Rounded Rectangle 17416', 'Right Arrow 17417', 'Rounded Rectangle 17418'):
        remove(s, nm)

    # interfaces
    L.text(sh, 0.80, 1.64, 2.30, 0.18, ['ROLE-SCOPED INTERFACES · React 19 + Vite'], size=7.5, bold=True, color=BLUE,
           align='l', anchor='m')
    tiles = [
        ('Ambulance', ':5174', 'crew intake, route', lambda x, y: L.ambulance(sh, x, y + 0.08, 0.20), None),
        ('Hospital', ':5175', 'referral answers, capacity', lambda x, y: L.hospital(sh, x + 0.06, y + 0.04, 0.28), None),
        ('Doctor', ':5177', 'clinical workspace', lambda x, y: L.doctor(sh, x + 0.04, y + 0.03, 0.30), 'dev'),
        ('Patient', ':5176', 'case status and updates', lambda x, y: L.patient(sh, x + 0.04, y + 0.03, 0.30), None),
        ('Management', ':5173', 'regional dashboard', lambda x, y: L.monitor(sh, x + 0.02, y + 0.05, 0.27), None),
    ]
    ty = 1.86
    for name, port, sub, icon, st in tiles:
        c = L.card(sh, 0.80, ty, 2.25, 0.37, radius=0.08, shadow=False, name=f'{name} interface')
        L.fill_text(c.text_frame, [[(name, {'size': 9, 'bold': True, 'color': NAVY}), (f'  {port}', {'size': 7.5, 'color': MUTED})],
                                   [(sub, {'size': 7.5, 'color': MUTED})]], anchor='m', align='l', ins=(0.52, 0.0, 0.04, 0.0))
        icon(0.86, ty)
        if st:
            L.status_pill(sh, 2.36, ty + 0.06, 0.62, 0.15, 'IN DEV', 'dev', size=6.5)
        ty += 0.41
    demo = L.card(sh, 0.80, 3.98, 2.25, 0.56, fill=AMBER_TINT, line=AMBER, lw=1.0, radius=0.08, shadow=False, dash=DASH,
                  name='Demo portal (separate)')
    L.fill_text(demo.text_frame, [[('Demo portal', {'size': 9, 'bold': True, 'color': AMBER_TX}), ('  :5178 · separate', {'size': 7.5, 'color': MUTED})],
                                  [('synthetic scenarios, same API and role checks; off in operational mode', {'size': 7, 'color': MUTED})]],
                anchor='m', align='l', ins=(0.52, 0.0, 0.04, 0.0))
    L.monitor(sh, 0.86, 4.10, 0.27, content='play', name='Demo portal illustration')

    # flows between interfaces and backend
    L.arrow(sh, 3.08, 2.50, 3.70, 2.50, color=BLUE, lw=1.75, name='Commands')
    L.text(sh, 3.06, 2.28, 0.66, 0.18, ['commands'], size=7, color=MUTED)
    L.arrow(sh, 3.70, 3.15, 3.08, 3.15, color=BLUE, lw=1.75, dash=DASH, name='Live events')
    L.text(sh, 3.06, 3.18, 0.66, 0.18, ['live events'], size=7, color=MUTED)
    L.arrow(sh, 3.08, 4.26, 3.70, 4.26, color=AMBER, lw=1.25, dash=DASH, name='Same API')
    L.text(sh, 3.06, 4.04, 0.66, 0.18, ['same API'], size=7, color=AMBER_TX)

    # backend
    L.card(sh, 3.72, 1.64, 5.86, 2.92, fill=SLATE_BG, line=BLUE, lw=1.25, radius=0.16, name='JIVA backend')
    L.text(sh, 3.88, 1.69, 5.6, 0.22, [[('JIVA BACKEND', {'bold': True, 'color': NAVY}),
                                        (' · local API: Node.js 20 + Express, REST + Socket.IO', {'color': MUTED})]], size=9.5,
           align='l', anchor='m')
    sec = L.card(sh, 3.86, 1.96, 5.58, 0.44, fill=WHITE, line=SLATE, lw=1.0, radius=0.08, shadow=False, name='Security band')
    L.fill_text(sec.text_frame, [[('SECURITY & PRIVACY  ', {'size': 8, 'bold': True, 'color': NAVY}),
                                  ('role checks on the server for every request · Zod schemas, strict payload filtering · '
                                   'idempotent, audited commands · policy hash · SHA-256-hashed decision audits', {'size': 7.5, 'color': MUTED})]],
                anchor='m', align='l', ins=(0.42, 0.02, 0.06, 0.02))
    L.shield(sh, 3.95, 2.02, 0.32)
    cols = [
        ('REFERRALS & ROUTING', BLUE, ['eligibility engine', 'time-limited referrals, expiry', 'destination & reroute',
                                       'SHADOW feasibility check (no authority)', 'ETA unknown, never faked']),
        ('CLINICAL RECORDS', AMBER, ['encounters, arrival, handover', 'vitals with provenance', 'documents (SHA-256 checked)',
                                     'prescriptions, reports, notes', 'clinical catalogue']),
        ('REAL-TIME EVENTS', BLUE, ['canonical event bus', 'state engines project state', 'content-free clinical events',
                                    'per-recipient copies (Socket.IO)', 'Bedrock: advisory text only']),
    ]
    fx = 3.86
    for title, col, items in cols:
        c = L.card(sh, fx, 2.48, 1.80, 1.44, fill=WHITE, line=col, lw=1.0, radius=0.08, shadow=False, name=f'Backend: {title}')
        paras = [[(title, {'size': 8, 'bold': True, 'color': col if col == BLUE else AMBER_TX})]]
        paras += [([('• ' + it, {'size': 7.5, 'color': TEXT})], {'space_before': 1.5}) for it in items]
        L.fill_text(c.text_frame, paras, anchor='t', align='l', ins=(0.08, 0.06, 0.04, 0.02))
        if col == AMBER:
            L.status_pill(sh, fx + 0.08, 3.68, 1.10, 0.16, 'IN DEVELOPMENT', 'dev', size=6.5)
        fx += 1.89
    st = L.card(sh, 3.86, 4.00, 5.58, 0.46, fill=WHITE, line=SLATE, lw=1.0, radius=0.08, shadow=False, name='State and event stores')
    L.fill_text(st.text_frame, [[('STATE & EVENT STORES  ', {'size': 8, 'bold': True, 'color': NAVY}),
                                 ('in memory locally — lost on restart · DynamoDB store defined in code, not deployed', {'size': 7.5, 'color': MUTED})]],
                anchor='m', align='l', ins=(0.42, 0.02, 0.06, 0.02))
    L.database(sh, 3.97, 4.06, 0.26, 0.34)

    # integrations
    L.text(sh, 9.98, 1.64, 2.55, 0.18, ['DATA & INTEGRATIONS'], size=7.5, bold=True, color=BLUE, align='l', anchor='m')
    ints = [
        ('Maps', 'MapLibre GL + OpenStreetMap tiles', 'neutral'),
        ('Routing', 'Valhalla → OSRM → labelled mock fallback', 'neutral'),
        ('Facility data', 'public listings; capacity unknown unless reported', 'neutral'),
        ('AI', 'Amazon Bedrock advisory summaries; mock locally', 'plan'),
        ('AWS', 'CDK stack synthesized — not deployed', 'cdk'),
    ]
    iy = 1.86
    for t, d, kind in ints:
        f, b, c = L.STATUS[kind]
        card = L.card(sh, 9.98, iy, 2.55, 0.46, fill=WHITE if kind == 'neutral' else f, line=b, lw=1.0, radius=0.08,
                      shadow=False, name=f'Integration: {t}')
        L.fill_text(card.text_frame, [[(t, {'size': 8, 'bold': True, 'color': NAVY if kind == 'neutral' else c})],
                                      [(d, {'size': 7, 'color': MUTED})]], anchor='m', align='l', ins=(0.08, 0.0, 0.04, 0.0))
        iy += 0.535
    L.arrow(sh, 9.60, 3.10, 9.96, 3.10, color=BLUE, lw=1.5, head='triangle', name='Integration calls')

    # roadmap row with status markers (implementation status)
    for nm, x in (('Rounded Rectangle 17421', 0.80), ('Rounded Rectangle 17422', 2.775), ('Rounded Rectangle 17423', 4.75),
                  ('Rounded Rectangle 17424', 6.725)):
        c = find(s, nm)
        c.left, c.width, c.top, c.height = E(x), E(1.90), E(5.18), E(0.76)
    newp = [(8.70, 'Phase 5: Clinical Care Layer', 'Handover, doctor workspace, documents, prescriptions', 'dev'),
            (10.675, 'Next: Production Path', 'Durable storage, AWS deployment, real identity', 'plan')]
    for x, t, d, kind in newp:
        f, b, c = L.STATUS[kind]
        card = L.card(sh, x, 5.18, 1.90, 0.76, fill=SLATE_BG if kind == 'dev' else WHITE, line=b, lw=1.0, radius=0.12,
                      dash=DASH if kind == 'plan' else None, name=t)
        L.fill_text(card.text_frame, [[(t, {'size': 10, 'bold': True, 'color': c})],
                                      [(d, {'size': 9, 'color': MUTED})]], ins=(0.06, 0.04, 0.06, 0.04))
    for x, kind in ((0.80, 'ok'), (2.775, 'ok'), (4.75, 'ok'), (6.725, 'ok'), (8.70, 'wait'), (10.675, 'expired')):
        if kind == 'expired':
            L.shp(sh, MSO_SHAPE.OVAL, x + 1.90 - 0.13, 5.18 - 0.08, 0.20, 0.20, fill=WHITE, line=GREY, lw=1.25, name='Planned marker')
        else:
            L.badge(sh, x + 1.90 - 0.03, 5.18 + 0.02, 0.20, kind if kind == 'ok' else 'busy')
    lg = [(8.05, 'ok', 'implemented'), (9.35, 'busy', 'in development'), (10.85, None, 'planned')]
    for x, kind, label in lg:
        if kind:
            L.badge(sh, x, 4.97, 0.15, kind)
        else:
            L.shp(sh, MSO_SHAPE.OVAL, x - 0.075, 4.97 - 0.075, 0.15, 0.15, fill=WHITE, line=GREY, lw=1.25)
        L.text(sh, x + 0.12, 4.88, 1.2, 0.18, [label], size=7.5, color=MUTED, align='l', anchor='m')


# ================================================================ NEW — the doctor workspace
def doctor_slide():
    s = clone_frame_slide(S[2], 'THE DOCTOR WORKSPACE')
    sh = s.shapes
    panel(sh, 0.60, 1.20, 12.13, 5.62, name='Doctor workspace panel')
    heading(sh, 0.85, 1.27, 8.25, 'CLINICAL LAYER: from ambulance intake to clinician-authorised care')
    L.status_pill(sh, 9.20, 1.29, 3.30, 0.24, 'WORKING PROTOTYPE · IN DEVELOPMENT', 'dev', size=8)

    rows = [
        ('Ambulance intake', 'symptoms, description, onset; vitals typed in', lambda x, y: L.ambulance(sh, x - 0.05, y + 0.07, 0.24)),
        ('Patient encounter', 'opens when the hospital records arrival', lambda x, y: L.encounter_card(sh, x, y + 0.04, 0.32)),
        ('Symptoms and vitals', 'each reading: time, source, quality flag', lambda x, y: L.clipboard(sh, x + 0.06, y, 0.38)),
        ('Clinical documents', 'scans and reports, typed and versioned', lambda x, y: L.xray(sh, x + 0.06, y, 0.38)),
        ('Doctor review', 'only the assigned doctor sees the record', lambda x, y: L.doctor(sh, x + 0.02, y - 0.01, 0.40)),
        ('Clinical assessment', "the doctor's own reports and notes, versioned", lambda x, y: L.doc_page(sh, x + 0.06, y, 0.38)),
        ('Prescription draft', 'catalogue product; dose and timing by the doctor', lambda x, y: L.rx_card(sh, x + 0.06, y, 0.38, 'draft')),
        ('Clinician authorisation', 'explicit attestation of the reviewed version', lambda x, y: L.rx_card(sh, x + 0.06, y, 0.38, 'authorised')),
    ]
    RX, RW, RH, RG, RY = 0.85, 4.45, 0.50, 0.13, 1.70
    mids = []
    for i, (t, d, icon) in enumerate(rows):
        y = RY + i * (RH + RG)
        c = L.card(sh, RX, y, RW, RH, radius=0.10, name=f'Workspace step {i + 1}: {t}',
                   fill='F0FDF4' if i == 7 else WHITE, line=GREEN if i == 7 else SLATE, lw=1.5 if i == 7 else 1.0)
        L.fill_text(c.text_frame, [[(t, {'size': 10, 'bold': True, 'color': NAVY})], [(d, {'size': 8, 'color': MUTED})]],
                    anchor='m', align='l', ins=(1.02, 0.0, 0.06, 0.0))
        L.num_badge(sh, RX + 0.20, y + RH / 2, 0.26, i + 1, size=9)
        icon(RX + 0.40, y + 0.06)
        mids.append(y + RH / 2)
        if i < 7:
            L.shp(sh, MSO_SHAPE.ISOSCELES_TRIANGLE, RX + RW / 2 - 0.07, y + RH + 0.025, 0.14, 0.08, fill=BLUE, rot=180,
                  name='Flow chevron')

    BX, BW = 5.80, 6.68
    # R1 — the crew's record
    r1 = L.card(sh, BX, 1.70, BW, 1.76, fill=WHITE, line=SLATE, radius=0.14, name='Pre-hospital record detail')
    L.fill_text(r1.text_frame, [
        [('PRE-HOSPITAL RECORD', {'size': 9, 'bold': True, 'color': NAVY}), ('  written by the attending crew', {'size': 8, 'color': MUTED})],
        ([('• Symptoms, free-text description and onset time', {'size': 8.5, 'color': TEXT})], {'space_before': 3}),
        ([('• Vital signs typed in with time, source and a quality flag; a reason for any not taken', {'size': 8.5, 'color': TEXT})], {'space_before': 2}),
        ([('• Goes to the assigned doctor and the patient only — never to dispatch or routing', {'size': 8.5, 'color': TEXT})], {'space_before': 2}),
        ([('• No device integration; a labelled simulator exists in demo mode only', {'size': 8.5, 'color': TEXT})], {'space_before': 2}),
    ], anchor='m', align='l', ins=(1.62, 0.06, 0.10, 0.06))
    L.ambulance(sh, BX + 0.18, 2.43, 0.36)
    L.clipboard(sh, BX + 0.88, 2.28, 0.62)
    L.arrow(sh, RX + RW + 0.02, mids[0], BX - 0.02, mids[0], color=BLUE, lw=1.25)
    L.arrow(sh, RX + RW + 0.02, mids[2], BX - 0.02, mids[2], color=BLUE, lw=1.25, head='triangle', tail=None)
    # R2 — documents
    r2 = L.card(sh, BX, 3.59, BW, 1.39, fill=WHITE, line=SLATE, radius=0.14, name='Clinical documents detail')
    L.text(sh, BX + 0.16, 3.64, 6.3, 0.22, [[('CLINICAL DOCUMENTS', {'bold': True, 'color': NAVY}),
                                            ('  typed uploads, verified by size, SHA-256 and file type; new versions keep the old one', {'size': 8, 'color': MUTED})]],
           size=9, align='l', anchor='m')
    th = 0.52
    groups = [
        ('X-ray / CT / MRI', [(L.xray, L.doc_width(th)), (L.ct_scan, L.doc_width(th))]),
        ('ECG reports', [(L.ecg, L.ecg_width(th))]),
        ('Ultrasound examples', [(L.ultrasound, L.doc_width(th))]),
        ('Laboratory reports', [(L.lab_report, L.doc_width(th))]),
    ]
    gx = BX + 0.30
    slot = (BW - 0.60) / 4
    for label, fns in groups:
        tot = sum(w for _, w in fns) + 0.06 * (len(fns) - 1)
        x0 = gx + (slot - tot) / 2
        for fn, w in fns:
            fn(sh, x0, 3.93, th)
            x0 += w + 0.06
        L.text(sh, gx, 3.93 + th + 0.03, slot, 0.18, [label], size=8, bold=True, color=NAVY)
        gx += slot
    L.text(sh, BX + 0.16, 4.70, BW - 0.32, 0.20, ['Never interpreted automatically · demo samples are synthetic and labelled “not this patient”'],
           size=7.5, italic=True, color=MUTED)
    L.arrow(sh, RX + RW + 0.02, mids[3], BX - 0.02, mids[3], color=BLUE, lw=1.25)
    # R3 — catalogue and safeguards
    r3 = L.card(sh, BX, 5.11, BW, 1.50, fill=AMBER_TINT, line=AMBER, lw=1.0, radius=0.14, name='Clinical catalogue detail')
    L.fill_text(r3.text_frame, [
        [('CLINICAL CATALOGUE (JSON / CSV)', {'size': 9, 'bold': True, 'color': AMBER_TX}), ('  illustrative, versioned', {'size': 8, 'color': MUTED})],
        ([('Medication products (name, strength, form — no doses) and the synthetic sample library. Administrators import '
           'with a preview and explicit confirmation; nothing is overwritten silently.', {'size': 8, 'color': TEXT})], {'space_before': 2}),
    ], anchor='t', align='l', ins=(0.16, 0.08, 0.12, 0.04))
    flow = [('JSON / CSV file', 'neutral'), ('Preview (changes, invalid, duplicates)', 'neutral'),
            ('Admin confirms', 'neutral'), ('Catalogue version +1', 'done')]
    fxp = BX + 0.16
    for i, (t, kind) in enumerate(flow):
        w = 0.16 + len(t) * (0.056 if kind == 'done' else 0.052)
        f, b, c = L.STATUS[kind]
        L.pill(sh, fxp, 5.86, w, 0.24, t, f, b, NAVY if kind == 'neutral' else c, size=7.5, bold=kind == 'done',
               name=f'Import step: {t}')
        if i < len(flow) - 1:
            L.arrow(sh, fxp + w + 0.03, 5.98, fxp + w + 0.17, 5.98, color=AMBER, lw=1.25, size='sm')
        fxp += w + 0.20
    chips = ['Never diagnoses', 'Never suggests a dose', 'Never authorises', 'Drafts never reach the patient']
    cx = BX + 0.16
    for t in chips:
        w = 0.20 + len(t) * 0.062
        L.pill(sh, cx, 6.22, w, 0.24, t, WHITE, AMBER, AMBER_TX, size=7.5, name=f'Safeguard: {t}')
        cx += w + 0.10
    L.arrow(sh, BX - 0.02, mids[6], RX + RW + 0.02, mids[6], color=AMBER, lw=1.25, name='Medication picker')
    L.text(sh, RX + RW + 0.04, mids[6] - 0.20, 0.42, 0.16, ['picker'], size=6.5, italic=True, color=AMBER_TX)
    return s


# ================================================================ NEW — AWS architecture with status distinctions
def aws_slide():
    s = clone_frame_slide(S[2], 'AWS ARCHITECTURE (DEFINED IN CDK)')
    sh = s.shapes
    panel(sh, 0.60, 1.20, 12.13, 5.62, name='AWS panel')
    heading(sh, 0.85, 1.27, 11.6, 'CLOUD TARGET: the local mesh mapped to serverless AWS — synthesized with CDK, never deployed')

    # clients (outside the account)
    L.monitor(sh, 0.92, 3.05, 0.50, content='lines', name='Web apps')
    L.text(sh, 0.76, 3.58, 0.95, 0.40, ['Web apps', 'React + Vite'], size=8, bold=True, color=NAVY)

    BL, BT, BR, BB = 1.95, 1.70, 9.55, 6.65
    L.card(sh, BL, BT, BR - BL, BB - BT, fill=WHITE, line=BLUE, lw=1.25, radius=0.18, shadow=False, dash=DASH,
           name='AWS account boundary')
    L.text(sh, BL + 0.15, BT + 0.04, 7.2, 0.22, [[('AWS ACCOUNT', {'bold': True, 'color': BLUE}),
                                                 ('  ·  one CDK stack  ·  cdk synth passes  ·  nothing deployed', {'color': MUTED})]],
           size=8.5, align='l', anchor='m')

    def svc(x, y, w, h, title, sub, local=None, name=None, fill='F4F8FC'):
        c = L.card(sh, x, y, w, h, fill=fill, line=BLUE, lw=1.0, radius=0.08, shadow=False, name=name or title)
        paras = [[(title, {'size': 8.5, 'bold': True, 'color': NAVY})], [(sub, {'size': 7, 'color': MUTED})]]
        if local:
            paras.append([('local: ' + local, {'size': 6.5, 'italic': True, 'color': GREEN_TX})])
        L.fill_text(c.text_frame, paras, anchor='m', align='l', ins=(0.08, 0.02, 0.04, 0.02))
        return c

    # ingest row (top-left), the event bus (centre), processing (right), delivery row (bottom-left)
    svc(2.12, 2.12, 1.52, 0.78, 'API Gateway · REST', 'POST /api/events', 'Express REST')
    svc(3.84, 2.12, 1.36, 0.78, 'Lambda · ingestion', 'validates, puts events')
    svc(2.12, 3.18, 1.52, 0.62, 'Amazon Cognito', 'user pool · role groups', name='Amazon Cognito')
    bus = L.card(sh, 5.40, 2.05, 1.02, 3.30, fill=NAVY, line=NAVY, lw=1.0, radius=0.10, shadow=False, name='Amazon EventBridge')
    L.fill_text(bus.text_frame, [[('Amazon', {'size': 8.5, 'bold': True, 'color': WHITE})],
                                 [('EventBridge', {'size': 8.5, 'bold': True, 'color': WHITE})],
                                 ([('jiva-healthcare-mesh bus · 6 rules', {'size': 7, 'color': 'D6E2EF'})], {'space_before': 3}),
                                 ([('local: in-process event bus', {'size': 6.5, 'italic': True, 'color': 'BBF7D0'})], {'space_before': 3})],
                ins=(0.06, 0.04, 0.06, 0.04))
    svc(6.62, 2.12, 1.40, 1.95, 'Lambda · processors ×5', 'emergency · hospital · ambulance · routing · AI sidecar', 'state engines')
    svc(8.18, 2.12, 1.25, 0.62, 'Amazon DynamoDB', 'single table · PITR', 'in-memory stores')
    svc(8.18, 2.84, 1.25, 0.50, 'Amazon S3', 'audit reports · private')
    svc(8.18, 3.44, 1.25, 0.62, 'Amazon Bedrock', 'advisory summaries only', 'mock templates')
    svc(6.62, 4.42, 2.81, 0.50, 'Amazon SQS · dead-letter queue', 'failed rule targets after 2 retries')
    svc(3.84, 4.50, 1.36, 0.78, 'Lambda · realtime fan-out', 'per-recipient delivery')
    svc(2.12, 4.50, 1.52, 0.78, 'API Gateway · WebSocket', 'IAM-authorised connect', 'Socket.IO')
    cw = L.card(sh, 2.12, 5.56, 7.31, 0.42, fill='F4F8FC', line=BLUE, lw=1.0, radius=0.08, shadow=False, name='Amazon CloudWatch')
    L.fill_text(cw.text_frame, [[('Amazon CloudWatch', {'size': 8.5, 'bold': True, 'color': NAVY}),
                                 ('   metrics, dashboard and two alarms for the operational event mesh', {'size': 7, 'color': MUTED})]],
                anchor='m', align='l', ins=(0.10, 0.0, 0.04, 0.0))
    note = L.card(sh, 2.12, 6.08, 7.31, 0.42, fill=AMBER_TINT, line=AMBER, lw=1.0, radius=0.08, shadow=False, dash=DASH,
                  name='Clinical records not in the stack')
    L.fill_text(note.text_frame, [[('Not in this stack: ', {'size': 7.5, 'bold': True, 'color': AMBER_TX}),
                                   ('clinical records, documents and the doctor workspace exist only in the local API today.',
                                    {'size': 7.5, 'color': TEXT})]], anchor='m', align='l', ins=(0.10, 0.0, 0.04, 0.0))

    # flows (straight, short, no crossings)
    L.arrow(sh, 1.62, 3.10, 2.10, 2.62, color=BLUE, lw=1.5, name='Commands in')
    L.arrow(sh, 3.66, 2.51, 3.82, 2.51, color=BLUE, lw=1.5)
    L.arrow(sh, 5.22, 2.51, 5.38, 2.51, color=BLUE, lw=1.5, name='Put events')
    L.arrow(sh, 2.88, 3.16, 2.88, 2.92, color=MUTED, lw=1.0, dash=DASH, name='Authorizer')
    L.text(sh, 2.94, 2.92, 0.80, 0.22, ['authorizer'], size=6.5, italic=True, color=MUTED, align='l', anchor='m')
    L.arrow(sh, 6.44, 3.10, 6.60, 3.10, color=BLUE, lw=1.5, name='Rules to processors')
    for y in (2.43, 3.09, 3.75):
        L.arrow(sh, 8.04, y, 8.16, y, color=BLUE, lw=1.25)
    L.arrow(sh, 6.44, 4.67, 6.60, 4.67, color=GREY, lw=1.25, name='Failed targets')
    L.arrow(sh, 5.38, 4.89, 5.22, 4.89, color=BLUE, lw=1.5, name='Realtime broadcast rule')
    L.arrow(sh, 3.82, 4.89, 3.66, 4.89, color=BLUE, lw=1.5, name='Push to sockets')
    L.arrow(sh, 2.10, 4.75, 1.50, 4.00, color=BLUE, lw=1.5, dash=DASH, name='Live events out')
    L.text(sh, 4.00, 4.12, 1.38, 0.32, ['realtime rule (never the', 'feasibility trace)'], size=6.5, italic=True,
           color=MUTED, align='r', anchor='b')

    # status column
    SX, SW = 9.78, 2.75
    stat = [
        ('done', '✓ IMPLEMENTED & VERIFIED LOCALLY', 'The same logic runs in the local API (green “local” notes); the 7 Lambda bundles build and load; cdk synth passes.', 1.18),
        ('cdk', 'DEFINED IN CODE (CDK)', 'Every AWS resource on this slide — synthesized, never deployed.', 0.92),
        ('plan', 'NOT TESTED AGAINST LIVE AWS', 'Cognito tokens (fixtures only) · DynamoDB store · WebSocket fan-out · Bedrock (failure path only).', 1.18),
        ('dev', 'PLANNED', 'Clinical records on AWS (APIs, tables, S3 document storage) · deployment and live end-to-end tests · monitoring of clinical paths.', 1.30),
    ]
    sy = 1.70
    for kind, t, d, h in stat:
        f, b, c = L.STATUS[kind]
        card = L.card(sh, SX, sy, SW, h, fill=f, line=b, lw=1.25, radius=0.12, shadow=False,
                      dash=DASH if kind == 'dev' else None, name=f'Status: {t}')
        L.fill_text(card.text_frame, [[(t, {'size': 8.5, 'bold': True, 'color': c})],
                                      ([(d, {'size': 8, 'color': TEXT})], {'space_before': 2})], anchor='t', align='l',
                    ins=(0.12, 0.08, 0.10, 0.04))
        sy += h + 0.12
    return s


# ================================================================ SLIDE 4 — feasibility: restrained header icons, accuracy fix
def slide4_feasibility():
    s = S[3]
    sh = s.shapes
    L.monitor(sh, 0.84, 1.42, 0.40, content='lines', name='Web overlay illustration')
    g = L.group(sh, 'Local to cloud illustration')
    L.monitor(g.shapes, 4.90, 1.50, 0.30, content='lines')
    L.cloud(g.shapes, 5.21, 1.40, 0.44, 0.27)
    L.coin(sh, 9.16, 1.42, 0.42)
    replace_in_runs(find(s, 'Rounded Rectangle 17410'), 'ECG IoT simulated via JSON',
                    'no device integration')


# ================================================================ SLIDE 5 — impact: stakeholder illustrations, accuracy fixes
def slide5_impact():
    s = S[4]
    sh = s.shapes
    L.patient(sh, 2.50 - 0.20, 1.36, 0.40)
    L.ambulance(sh, 6.45 - L.ambulance_width(0.32) / 2, 1.38, 0.32)
    L.hospital(sh, 2.50 - L.hospital_width(0.36) / 2, 4.09, 0.36)
    L.monitor(sh, 6.45 - L.monitor_width(0.36) / 2, 4.10, 0.36)
    replace_in_runs(find(s, 'Rounded Rectangle 17411'), 'Guaranteed confirmed hospital bed before ambulance leaves the scene.',
                    'Destination set only after a hospital explicitly accepts, never assumed from a listing.')
    replace_in_runs(find(s, 'Rounded Rectangle 17412'), 'Structured pre-arrival vitals enable advance trauma & cath lab preparation.',
                    'Structured pre-arrival referral (required care, ETA) enables advance trauma & cath lab preparation.')
    imp = find(s, 'Rounded Rectangle 17414')
    replace_in_runs(imp, 'Dramatically reduces', 'Aims to reduce')
    replace_in_runs(imp, 'by eliminating secondary hospital transfers.', 'by avoiding secondary hospital transfers.')
    replace_in_runs(imp, '(Saves ~1.9M brain neurons per minute saved in ischemic stroke)',
                    '(Untreated ischaemic stroke destroys ~1.9M neurons per minute — Saver, Stroke 2006)')
    gov = find(s, 'Rounded Rectangle 17416')
    replace_in_runs(gov, 'Tamper-Evident SHA-256 Ledger', 'SHA-256 Decision Audit Hashes')
    replace_in_runs(gov, 'Cryptographic policy hashing (hashPolicy) and immutable SHA-256 event audit trails provide forensic accountability for clinical governance and medico-legal auditability.',
                    'Policy hashing (hashPolicy) and SHA-256-hashed routing-decision audits support governance review; a durable, tamper-evident ledger is planned.')


# ================================================================ SLIDE 6 — references: keep text; correct the app count
def slide6_references():
    s = S[5]
    replace_in_runs(find(s, 'Rounded Rectangle 17414'), '4 web apps', '5 role-scoped web apps + a separate demo portal')


# ---------------------------------------------------------------- build
slide1_title()
slide2_solution()
j = journey_slide()
r = responses_slide()
slide3_architecture()
d = doctor_slide()
a = aws_slide()
slide4_feasibility()
slide5_impact()
slide6_references()
move_slide(j, 2)
move_slide(r, 3)
move_slide(d, 5)
move_slide(a, 6)
prs.save(OUT)
print('saved', OUT, 'slides:', len(prs.slides))
