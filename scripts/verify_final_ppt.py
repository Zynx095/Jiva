import os
import pptx
from pptx import Presentation

pptx_path = 'ppt/JIVA_IDE_BOOTCAMP_FORMAT_FINAL.pptx'
pdf_path = 'ppt/JIVA_IDE_BOOTCAMP_FORMAT_FINAL.pdf'

print(f"PPTX Size: {os.path.getsize(pptx_path)} bytes")
print(f"PDF Size: {os.path.getsize(pdf_path)} bytes")

prs = Presentation(pptx_path)
print(f"Slide Dimensions: {prs.slide_width / 914400:.3f} x {prs.slide_height / 914400:.3f} inches")
print(f"Total Slides: {len(prs.slides)}")

for idx, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {idx + 1} ====================")
    for s in slide.shapes:
        pos = f"L={s.left/914400:.2f}, T={s.top/914400:.2f}, W={s.width/914400:.2f}, H={s.height/914400:.2f}"
        if s.has_text_frame:
            p_count = len(s.text_frame.paragraphs)
            char_count = sum(len(p.text) for p in s.text_frame.paragraphs)
            first_text = s.text_frame.paragraphs[0].text[:40].replace('\n', ' ') if p_count > 0 else ''
            print(f"  Shape: {s.name:25} | {pos} | Paras: {p_count:2} | Chars: {char_count:4} | '{first_text}...'")
        else:
            print(f"  Shape: {s.name:25} | {pos} | [No text frame]")
